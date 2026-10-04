"""Saved bytes, extraction limits and layout claims must stay distinguishable."""

import hashlib
import io
import zipfile

import pytest

from rune.capabilities.document_inspection import inspect_bytes, inspect_document


def test_docx_keeps_body_order_and_does_not_invent_pagination(tmp_path, monkeypatch):
    from docx import Document

    from rune.agent.loop import NativeAgentLoop

    monkeypatch.setenv("RUNE_HOME", str(tmp_path / "state"))
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Mina"
    doc.add_paragraph("After")
    path = tmp_path / "brief.docx"
    doc.save(path)
    result = inspect_document(path)
    assert result["text"].index("Before") < result["text"].index("Mina") < result["text"].index("After")
    assert result["facts"]["page_count"] is None and not result["visual_verified"]
    assert result["facts"]["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    agent = NativeAgentLoop()
    agent._files_written.add(str(path))
    assert "Mina" in agent._gather_artifact()


def test_workbook_reports_formula_instead_of_an_empty_verified_value():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append([2, 3, "=SUM(A1:B1)"])
    stream = io.BytesIO()
    wb.save(stream)
    result = inspect_bytes(stream.getvalue(), "xlsx")
    assert "=SUM(A1:B1)" in result["text"]
    assert result["facts"]["formulas_without_cached_values"] == 1
    assert any("not recalculated" in w for w in result["warnings"])


def test_pdf_page_count_is_measured_and_empty_pages_need_inspection():
    from pypdf import PdfWriter

    pdf = PdfWriter()
    for _ in range(2):
        pdf.add_blank_page(width=600, height=800)
    stream = io.BytesIO()
    pdf.write(stream)
    result = inspect_bytes(stream.getvalue(), "pdf")
    assert result["facts"]["page_count"] == 2
    assert result["facts"]["pages_without_text"] == [1, 2]
    assert not result["visual_verified"]


def test_slide_outside_canvas_is_not_silently_accepted():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(prs.slide_width - 10, 0, 200, 100)
    shape.text = "Outside"
    stream = io.BytesIO()
    prs.save(stream)
    result = inspect_bytes(stream.getvalue(), "pptx")
    assert result["facts"]["shapes_outside_slide"] == [{"slide": 1, "shape": shape.shape_id}]


def test_truncation_and_zip_expansion_remain_unverified():
    result = inspect_bytes(b"123456789", "txt", 4)
    assert result["text"] == "1234" and result["truncated"]
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("oversized.xml", b"x" * 50_000_001)
    with pytest.raises(ValueError, match="Expanded"):
        inspect_bytes(stream.getvalue(), "docx")


def test_inspection_honors_isolation_boundary(tmp_path, monkeypatch):
    root = tmp_path / "work"
    root.mkdir()
    outside = tmp_path / "private.txt"
    outside.write_text("private")
    monkeypatch.setenv("RUNE_ISOLATION_ROOT", str(root))
    with pytest.raises(ValueError):
        inspect_document(outside)
