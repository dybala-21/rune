"""Check actual table errors, immutable requirements, recovery and completion evidence."""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest

from rune.agent.table_acceptance import TableAcceptance, acceptance_scope, extract_plan
from rune.api.trust import build_trust_payload
from rune.capabilities.table_acceptance import (
    TablePlan,
    compare_table,
    expected_table,
    read_tabular,
)

REQUEST = "Exclude cancelled orders, remove exact duplicates, sum amount by team. Save summary.csv."
SOURCE = ("id,team,status,amount\n1,A,confirmed,10.005\n1,A,confirmed,10.005\n"
          "2,A,cancelled,99\n3,B,confirmed,5.005\n4,A,confirmed,2.000\n")


def plan(**changes):
    return TablePlan.model_validate({
        "applicable": True, "requirements": [REQUEST],
        "filters": [{"column": "status", "operator": "ne", "values": ["cancelled"]}],
        "deduplicate_by": ["id", "team", "status", "amount"], "group_by": ["team"],
        "aggregates": [{"name": "amount", "column": "amount", "operation": "sum", "decimals": 2}],
        "output_names": ["summary.csv"], **changes,
    })


@pytest.fixture
def office(tmp_path, monkeypatch):
    source, output = tmp_path / "sales.csv", tmp_path / "summary.csv"
    source.write_text(SOURCE)
    output.write_text("team,amount\nB,5.01\nA,12.01\n")
    guard = SimpleNamespace(validate_file_read_path=lambda path: SimpleNamespace(
        allowed=path.startswith(str(tmp_path) + "/"), reason="outside fixture workspace"))
    monkeypatch.setattr("rune.safety.guardian.get_guardian", lambda: guard)
    calls = []

    async def extract(*args):
        calls.append(args)
        return plan()

    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", extract)
    return source, output, calls


@pytest.mark.asyncio
async def test_real_output_can_be_repaired_without_changing_contract(office):
    source, output, calls = office
    state = TableAcceptance(REQUEST, required=True)
    assert await state.blocker()
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    assert await state.blocker()
    output.write_text("team,amount\nA,17.02\n")  # Correct grand total, wrong grouping and coverage.
    failed = await state.verify(contract["id"], str(output), None, 1)
    assert not failed.success
    assert failed.metadata["tableVerification"]["status"] == "fail"
    assert await state.blocker()
    output.write_text("team,amount\nB,5.01\nA,12.01\n")
    passed = await state.verify(contract["id"], str(output), None, 1)
    assert passed.success, passed.output
    assert passed.metadata["tableVerification"]["stats"] == {
        "source_rows": 5, "filtered_rows": 1, "duplicates_removed": 1, "selected_rows": 3, "output_rows": 2,
    }
    assert await state.blocker() is None
    assert (await state.requirements(str(source), None)).metadata["tableContract"] == contract
    assert len(calls) == 1
    trust = build_trust_payload(SimpleNamespace(reason="completed", table_acceptance=state.snapshot()))
    assert trust["verified"] and trust["verificationRequired"]
    output.write_text("team,amount\nA,17.02\n")
    assert await state.blocker()
    trust = build_trust_payload(SimpleNamespace(reason="completed", mech_check="pass", table_acceptance=state.snapshot()))
    assert not trust["verified"] and trust["verificationStatus"] == "inconclusive"


async def test_missing_table_checks_select_recovery_tools_and_release_repairs(office):
    source, output, _ = office
    state = TableAcceptance(REQUEST, required=True)
    assert state.recovery_tools() is None
    assert await state.blocker()
    assert "table_requirements" in state.recovery_tools()
    assert "bash_execute" not in state.recovery_tools()
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    assert state.recovery_tools() is None
    assert await state.blocker()
    assert "table_verify" in state.recovery_tools()
    output.write_text("team,amount\nA,999\n")
    assert not (await state.verify(contract["id"], str(output), None, 1)).success
    assert state.recovery_tools() is None
    output.write_text("team,amount\nA,12.01\nB,5.01\n")
    assert (await state.verify(contract["id"], str(output), None, 1)).success
    assert await state.blocker() is None


@pytest.mark.parametrize("content", [
    "team,amount\nA,12.01\nB,5.01\nB,5.01\n",  # extra duplicate
    "team,amount\nA,12.00\nB,5.01\n",  # wrong rounding
    "team,amount\nA,5.01\nB,12.01\n",  # same total, swapped groups
    "team,amount,extra\nA,12.01,x\nB,5.01,x\n",
])
def test_wrong_tables_fail_even_when_totals_look_plausible(content):
    headers, source = read_tabular(SOURCE.encode(), ".csv")
    expected, _ = expected_table(plan(), headers, source)
    headers, actual = read_tabular(content.encode(), ".csv")
    assert compare_table(plan(), expected, headers, actual)["status"] == "fail"


@pytest.mark.parametrize("operation,total", [("sum", "17.01"), ("mean", "5.67"), ("count", "3.00")])
def test_grand_total_uses_selected_source_rows_before_group_rounding(operation, total):
    contract = plan(grand_total=["전체 합계"], aggregates=[{
        "name": "amount", "operation": operation, "column": "amount", "decimals": 2}])
    headers, source = read_tabular(SOURCE.encode(), ".csv")
    expected, stats = expected_table(contract, headers, source)
    assert expected[-1] == ("전체 합계", total)
    assert stats["selected_rows"] == 3 and stats["output_rows"] == 3
    columns = ["team", "amount"]
    output = [dict(zip(columns, row, strict=True)) for row in expected]
    assert compare_table(contract, expected, columns, output)["status"] == "pass"
    for wrong in (output[:-1], output + [output[-1]], output[:-1] + [{"team": "전체 합계", "amount": "999"}]):
        assert compare_table(contract, expected, columns, wrong)["status"] == "fail"


def test_grand_total_labels_cannot_hide_a_source_group():
    headers, rows = read_tabular(SOURCE.encode(), ".csv")
    with pytest.raises(ValueError, match="collides"):
        expected_table(plan(grand_total=["A"]), headers, rows)
    with pytest.raises(ValueError, match="grouping column"):
        plan(grand_total=["Total", ""])
    with pytest.raises(ValueError):
        plan(applicable=False, unverified=["unsupported"], filters=[], deduplicate_by=[],
             group_by=[], aggregates=[], grand_total=["Total"])


def test_unspecified_display_labels_do_not_become_data_requirements():
    contract = plan(grand_total=["Total"], exact_headers=False, exact_total_label=False)
    headers, source = read_tabular(SOURCE.encode(), ".csv")
    expected, _ = expected_table(contract, headers, source)
    columns, rows = read_tabular(b"Department,Total amount\nA,12.01\nB,5.01\nGrand total,17.01\n", ".csv")
    assert compare_table(contract, expected, columns, rows)["status"] == "pass"
    assert compare_table(contract.model_copy(update={"exact_headers": True}), expected, columns, rows)["status"] == "fail"
    assert compare_table(contract.model_copy(update={"exact_total_label": True}), expected, columns, rows)["status"] == "fail"
    for wrong in (rows[:-1], rows + [rows[-1]], rows[1:], rows[:-1] + [{"Department": "Grand total", "Total amount": "17.02"}]):
        assert compare_table(contract, expected, columns, wrong)["status"] == "fail"


@pytest.mark.parametrize("bom", [True, False])
async def test_requested_csv_bom_is_verified_from_file_bytes(office, monkeypatch, bom):
    source, output, _ = office

    async def extract(*args):
        return plan(csv_bom=bom)

    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", extract)
    state = TableAcceptance(REQUEST, required=True)
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    content = "team,amount\nA,12.01\nB,5.01\n"
    output.write_text(content, encoding="utf-8" if bom else "utf-8-sig")
    bad = await state.verify(contract["id"], str(output), None, 1)
    assert not bad.success and '"check": "csv_bom"' in bad.output
    output.write_text(content, encoding="utf-8-sig" if bom else "utf-8")
    good = await state.verify(contract["id"], str(output), None, 1)
    assert good.success and "csv_bom" in good.metadata["tableVerification"]["checks"]
    assert await state.blocker() is None


def test_conflicting_duplicate_keys_are_not_silently_dropped():
    headers, rows = read_tabular(b"id,team,status,amount\n1,A,confirmed,5\n1,B,confirmed,9\n", ".csv")
    with pytest.raises(ValueError, match="Conflicting rows"):
        expected_table(plan(deduplicate_by=["id"]), headers, rows)


def test_empty_selection_and_nonfinite_numbers():
    headers, rows = read_tabular(b"id,team,status,amount\n1,A,cancelled,5\n", ".csv")
    expected, stats = expected_table(plan(), headers, rows)
    assert expected == [] and stats["selected_rows"] == 0
    headers, output = read_tabular(b"team,amount\n", ".csv")
    assert compare_table(plan(), expected, headers, output)["status"] == "pass"
    with pytest.raises(ValueError, match="non-finite"):
        compare_table(plan(), [("A", "1")], headers, [{"team": "A", "amount": "NaN"}])


def test_grouped_amounts_are_verified_without_changing_source_values():
    data = b'id,team,status,amount\n1,A,confirmed,"12,500.50"\n2,A,confirmed,"-1,000.25"\n3,B,confirmed,0\n'
    headers, rows = read_tabular(data, ".csv")
    expected, _ = expected_table(plan(grand_total=["Total"]), headers, rows)
    assert expected == [("A", "11500.25"), ("B", "0.00"), ("Total", "11500.25")]
    assert rows[0]["amount"] == "12,500.50"
    columns, output = read_tabular(b'team,amount\nA,"11,500.25"\nB,0\nTotal,"11,500.25"\n', ".csv")
    assert compare_table(plan(grand_total=["Total"]), expected, columns, output)["status"] == "pass"


@pytest.mark.parametrize("value", ["12,34", "1,234,56", "1.234,56", "1,,234", "1,234_000", "", None, True, "NaN", "Infinity"])
def test_invalid_amounts_never_become_valid_totals(value):
    from rune.capabilities.table_acceptance import number

    with pytest.raises(ValueError):
        number(value)


@pytest.mark.asyncio
async def test_xlsx_readback_rejects_formulas_and_requires_sheet(office, monkeypatch):
    from openpyxl import Workbook

    source, output, _ = office
    async def xlsx_plan(*args):
        return plan(output_names=["summary.xlsx"], output_sheet="Summary")
    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", xlsx_plan)
    state = TableAcceptance(REQUEST)
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    output = output.with_suffix(".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["team", "amount"])
    ws.append(["A", 12.01])
    ws.append(["B", 5.01])
    wb.create_sheet("Notes")
    wb.save(output)
    result = await state.verify(contract["id"], str(output), "Notes", 1)
    assert not result.success and "worksheet" in result.output
    assert (await state.verify(contract["id"], str(output), None, 1)).success
    assert (await state.verify(contract["id"], str(output), "Summary", 1)).success
    ws["B2"] = "12.01"
    wb.save(output)
    result = await state.verify(contract["id"], str(output), "Summary", 1)
    assert not result.success and "numeric_cells" in result.output
    ws["B2"] = "=12.01"
    wb.save(output)
    result = await state.verify(contract["id"], str(output), "Summary", 1)
    assert not result.success and "Formula" in result.output
    wb.close()


@pytest.mark.asyncio
async def test_source_revision_and_read_policy_cannot_be_bypassed(office, tmp_path):
    source, output, _ = office
    state = TableAcceptance(REQUEST)
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    with pytest.raises(ValueError, match="outside fixture"):
        await state.requirements("/etc/hosts", None)
    source.write_text(SOURCE.replace("10.005", "99"))
    with pytest.raises(ValueError, match="Source changed"):
        await state.requirements(str(source), None)
    result = await state.verify(contract["id"], str(output), None, 1)
    assert not result.success and "Source changed" in result.output


@pytest.mark.asyncio
async def test_durable_contract_survives_restart_but_checks_must_run_again(office, tmp_path, monkeypatch):
    from rune.agent.execution_journal import ExecutionJournal, journal_scope, reconcile
    from rune.agent.tool_adapter import ToolAdapterOptions, build_tool_set
    from rune.api.run_snapshot import RunSnapshots
    from rune.api.run_store import RunStore
    from rune.capabilities.registry import CapabilityRegistry
    from rune.capabilities.table_checks import register_table_checks

    source, output, calls = office
    store = RunStore(tmp_path / "runs.sqlite")
    runs = RunSnapshots(store)
    runs.start("first", "session", REQUEST)
    registry = CapabilityRegistry()
    register_table_checks(registry)
    state = TableAcceptance(REQUEST)
    tools = build_tool_set(ToolAdapterOptions(enable_guardian=False, workspace_root=str(tmp_path),
                                             table_acceptance=state), registry=registry)
    with journal_scope(ExecutionJournal(store, "first", str(tmp_path))):
        result = json.loads(await tools["table_requirements"].function(source_path=source.name))
        contract = result["contract"]
        verified = json.loads(await tools["table_verify"].function(contract_id=contract["id"], output_path=output.name))
        assert verified["status"] == "pass"
    records = reconcile(store.attempts("first"))
    store.close()
    store = RunStore(tmp_path / "runs.sqlite")
    store.open()
    records = reconcile(store.attempts("first"))
    restored = TableAcceptance(REQUEST, previous=records)
    assert await restored.blocker()  # A prior passing check is not fresh evidence.
    assert (await restored.requirements(str(source), None)).metadata["tableContract"] == contract
    assert len(calls) == 1
    assert (await restored.verify(contract["id"], str(output), None, 1)).success
    assert await restored.blocker() is None
    unrelated = TableAcceptance("A different request", previous=records)
    assert not unrelated.contracts
    store.close()


@pytest.mark.asyncio
async def test_unavailable_extractor_and_unsupported_scope_never_pass(office, monkeypatch):
    from rune.capabilities.table_checks import TableRequirementsParams, table_requirements

    source, _, _ = office
    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", extract_plan)

    async def unavailable(*args):
        return None

    monkeypatch.setattr("rune.agent.requirement_gate._completion", unavailable)
    state = TableAcceptance(REQUEST, required=True)
    with acceptance_scope(state):
        result = await table_requirements(TableRequirementsParams(source_path=str(source)))
    assert not result.success and await state.blocker()
    assert not build_trust_payload(SimpleNamespace(reason="completed", mech_check="pass", table_acceptance=state.snapshot()))["verified"]


@pytest.mark.asyncio
async def test_completion_gate_blocks_unchecked_table_even_with_passing_code_tests():
    from rune.agent.loop import NativeAgentLoop

    loop = NativeAgentLoop()
    loop._requires_execution = True
    loop._table_acceptance = TableAcceptance(REQUEST, required=True)
    ok, messages, blocks = await loop._finalize_gates([], 0)
    assert not ok and blocks == 1 and messages
    assert loop._completion_check["name"] == "Table requirements"


def test_table_failure_dominates_other_passing_checks():
    trace = SimpleNamespace(reason="completed", mech_check="pass",
                            table_acceptance={"required": True, "status": "fail"})
    payload = build_trust_payload(trace)
    assert payload["verificationStatus"] == "failed" and not payload["verified"]


@pytest.mark.asyncio
async def test_unsupported_data_is_distinct_from_layout_and_returned_contract_is_a_copy(office, monkeypatch):
    source, output, _ = office

    async def extract(*args):
        return plan(out_of_scope=["Use company report styling"])

    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", extract)
    state = TableAcceptance(REQUEST)
    response = await state.requirements(str(source), None)
    contract = response.metadata["tableContract"]
    contract["plan"]["filters"] = []
    assert (await state.verify(contract["id"], str(output), None, 1)).success
    assert await state.blocker() is None
    assert state.snapshot()["status"] == "pass"
    assert state.snapshot()["out_of_scope"] == ["Use company report styling"]

    async def ambiguous(*args):
        return plan(unverified=["Convert mixed currencies without a specified exchange rate"])

    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", ambiguous)
    state = TableAcceptance(REQUEST)
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    assert (await state.verify(contract["id"], str(output), None, 1)).success
    assert await state.blocker() is None
    assert state.snapshot()["status"] == "inconclusive"


@pytest.mark.asyncio
async def test_source_alias_retarget_cannot_change_a_fixed_contract(office):
    source, output, _ = office
    alias = source.with_name("input.csv")
    other = source.with_name("other.csv")
    other.write_bytes(source.read_bytes())
    alias.symlink_to(source)
    state = TableAcceptance(REQUEST)
    contract = (await state.requirements(str(alias), None)).metadata["tableContract"]
    alias.unlink()
    alias.symlink_to(other)
    with pytest.raises(ValueError, match="Source changed"):
        await state.requirements(str(alias), None)
    assert not (await state.verify(contract["id"], str(output), None, 1)).success


@pytest.mark.asyncio
async def test_values_only_xlsx_source_uses_numeric_cells(office):
    import csv

    from openpyxl import Workbook

    source, output, _ = office
    workbook = Workbook()
    for index, row in enumerate(csv.reader(SOURCE.splitlines())):
        if index:
            row[-1] = float(row[-1])
        workbook.active.append(row)
    source = source.with_suffix(".xlsx")
    workbook.save(source)
    workbook.close()
    state = TableAcceptance(REQUEST)
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    assert (await state.verify(contract["id"], str(output), None, 1)).success
