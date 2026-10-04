"""Rendering supplies a real page image, not a fabricated layout verdict."""

import hashlib
import io
import shutil
import zipfile
from unittest.mock import AsyncMock

import pytest

from rune.capabilities.document_preview import (
    DocumentPreviewParams,
    _check_package,
    document_preview,
)
from rune.safety.execution_environment import environment_scope
from rune.safety.verification import CheckResult


def test_external_resources_and_macros_are_refused():
    for name, data in (("word/vbaProject.bin", b"macro"),
                       ("word/_rels/document.xml.rels", b'<Relationships><Relationship TargetMode="External" Type="image" Target="https://example.org/private"/></Relationships>')):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr(name, data)
        with pytest.raises(ValueError):
            _check_package(stream.getvalue())


@pytest.mark.asyncio
async def test_blocked_renderer_never_claims_a_preview(tmp_path, monkeypatch):
    from pypdf import PdfWriter

    pdf = PdfWriter()
    pdf.add_blank_page(width=600, height=800)
    path = tmp_path / "one.pdf"
    pdf.write(path)
    blocked = AsyncMock(return_value=CheckResult(error="Execution policy denied"))
    monkeypatch.setattr("rune.capabilities.document_preview.run_check", blocked)
    with environment_scope(str(tmp_path)):
        result = await document_preview(DocumentPreviewParams(path=str(path)))
    assert not result.success and "denied" in result.error
    assert not result.metadata.get("image_base64")
    assert not list(tmp_path.glob(".rune-preview-*"))


@pytest.mark.asyncio
@pytest.mark.skipif(not shutil.which("soffice") or not shutil.which("pdftoppm"), reason="Headless renderers not installed")
async def test_real_word_preview_keeps_source_and_delivers_pixels(tmp_path, monkeypatch):
    from docx import Document

    from rune.agent.tool_output import ToolOutput, output_for_model
    from rune.safety.resource_locks import capability_access

    monkeypatch.setenv("RUNE_HOME", str(tmp_path / "home"))
    doc = Document()
    doc.add_heading("Release plan", 0)
    doc.add_paragraph("Owner Mina. Deadline Monday.")
    source = tmp_path / "brief.docx"
    doc.save(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    with environment_scope(str(tmp_path)), capability_access("document_preview", {"path": str(source)}):
        result = await document_preview(DocumentPreviewParams(path=str(source)))
    assert result.success, result.error
    preview = result.metadata["preview"]
    assert preview["source_sha256"] == digest and preview["page_count"] == 1
    assert not preview["visual_verified"]
    output = output_for_model(result.output, "document_preview", result)
    assert isinstance(output, ToolOutput) and len(output.images) == 1
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert not list(tmp_path.glob(".rune-preview-*"))
