"""Regressions from the live office and coding workflows."""

import json
from unittest.mock import AsyncMock

import pytest

from rune.agent.litellm_adapter import StreamResult, _is_test_command
from rune.agent.provenance import ArtifactLedger


def stream(root, request="read office/source.xlsx", **kwargs):
    return StreamResult(
        model="openai/gpt-6-astra", workspace_root=str(root), request=request,
        messages=[{"role": "user", "content": "create old.pptx from old.csv"},
                  {"role": "assistant", "content": "done"},
                  {"role": "user", "content": request}],
        tool_schemas=[], tool_lookup={}, max_tokens=4096, temperature=0,
        request_tokens_limit=10000, response_tokens_limit=4096, **kwargs,
    )


def test_document_read_uses_selected_workspace_and_current_request(tmp_path, monkeypatch):
    server, workspace = tmp_path / "server", tmp_path / "workspace"
    server.mkdir()
    (workspace / "office").mkdir(parents=True)
    monkeypatch.chdir(server)
    (workspace / "office/source.xlsx").write_bytes(b"fixture")
    result = stream(workspace)
    result._record_provenance("document_read", {"path": "office/source.xlsx"}, "Workbook contents")
    assert result._ledger().referenced == {"source.xlsx"}
    assert result._unresolved_artifacts() == []
    assert result._ledger().read_ok == {"source.xlsx"}
    result._record_provenance("document_read", {"path": "office/missing.xlsx"}, "Error: missing")
    assert "missing.xlsx" not in result._ledger().read_ok


def test_failed_read_of_existing_file_is_not_success_or_absence(tmp_path):
    (tmp_path / "source.xlsx").write_bytes(b"unreadable")
    result = stream(tmp_path, "read source.xlsx")
    result._record_provenance("document_read", {"path": "source.xlsx"}, "Error: unsupported encryption")
    assert result._unresolved_artifacts() == ["source.xlsx"]
    assert not result._ledger().known_absent


@pytest.mark.parametrize("role", ["input", "preserve", None])
def test_hashing_an_existing_file_does_not_report_it_missing(tmp_path, role):
    (tmp_path / "coding").mkdir()
    (tmp_path / "coding/expenses.csv").write_text("amount\n12.30\n")
    result = stream(tmp_path, "옵션을 추가하고 coding/expenses.csv는 수정하지 마")
    if role:
        result._ledger().roles["expenses.csv"] = role
    result._record_provenance("bash_execute", {"command": "shasum coding/expenses.csv"}, "hash")
    assert result._unresolved_artifacts() == []
    assert not result._ledger().read_ok


def test_preserved_file_is_not_a_required_read_or_a_new_output(tmp_path):
    from rune.agent.postconditions import check, derive

    source = tmp_path / "expenses.csv"
    source.write_text("amount\n12.30\n")
    ledger = ArtifactLedger.for_request("leave expenses.csv unchanged", str(tmp_path))
    ledger.roles["expenses.csv"] = "preserve"
    conditions = derive(ledger.roles, tmp_path)
    source.unlink()
    ledger.record_read("expenses.csv", False)
    assert ledger.unresolved() == []
    assert ledger.is_phantom("expenses.csv")
    assert check(conditions, tmp_path)


async def test_command_suggestions_do_not_repeat_on_each_evaluation(monkeypatch):
    from types import SimpleNamespace

    from rune.proactive.engine import ProactiveEngine
    from rune.proactive.prediction.types import PredictionResult

    result = PredictionResult(tool_predictions=[("bash:unittest", .7), ("bash:<<'PY'", .7)])
    monkeypatch.setattr("rune.proactive.prediction.engine.get_prediction_engine", lambda: SimpleNamespace(predict=lambda context: result))
    engine = ProactiveEngine()
    monkeypatch.setattr(ProactiveEngine, "_gather_context", AsyncMock(return_value={}))
    first = await engine.evaluate({})
    assert [s.title for s in first if s.source == "behavior_prediction"] == ["Run unittest?"]
    assert not [s for s in await engine.evaluate({}) if s.source == "behavior_prediction"]


def test_missing_input_feedback_keeps_the_real_request_for_language_context():
    from rune.agent.provenance import unresolved_stop_note

    request = "coding/정산규정.md에 적힌 예외 규칙만 비교해 줘"
    note = unresolved_stop_note(["정산규정.md"], request)
    assert request in note
    assert "not a new user request" in note


def test_legacy_slides_still_receive_content_verification(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from rune.capabilities.document import DocBlock, DocumentCreateParams
    from rune.capabilities.document_bundle import _verify_file

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(3))
    box.text = "기존 보고서\n합계 100원"
    path = tmp_path / "old.pptx"
    prs.save(str(path))
    spec = DocumentCreateParams(path=str(path), format="pptx", title="기존 보고서",
                                blocks=[DocBlock(type="paragraph", text="합계 100원")])
    _verify_file(path, spec)
    box.text = "기존 보고서\n합계 200원"
    prs.save(str(path))
    with pytest.raises(ValueError, match="lost supplied text"):
        _verify_file(path, spec)


def test_same_basename_in_another_subdirectory_is_not_evidence(tmp_path):
    ledger = ArtifactLedger.for_request("read office/source.xlsx", str(tmp_path))
    ledger.record_read("office/source.xlsx", False)
    ledger.record_read("other/source.xlsx", True)
    assert ledger.unresolved() == ["source.xlsx"]
    ledger.record_read("office/source.xlsx", True)
    assert ledger.unresolved() == []


async def test_missing_input_does_not_trigger_parent_directory_search(tmp_path, monkeypatch):
    root = tmp_path / "project"
    root.mkdir()
    tool = AsyncMock(return_value="private listing")
    result = stream(root, "coding/정산규정.md에 적힌 규칙만 검토해")
    result._tool_lookup["file_list"] = tool
    monkeypatch.setattr("rune.agent.provenance.classify_roles", AsyncMock(return_value={"정산규정.md": "input"}))
    answer = await result._execute_tool("file_list", {"path": str(tmp_path)})
    assert answer.startswith("BLOCKED")
    tool.assert_not_awaited()
    assert result._unresolved_artifacts() == ["정산규정.md"]


@pytest.mark.parametrize("command", [
    "python3 -B -m unittest discover -s tests -v",
    "python3 -I -m pytest", "./.venv/bin/python -m pytest",
    "uv run python3 -W error -m unittest", 'cd "a b" && python3.12 -X dev -m pytest',
    "FOO=bar env PYTHONDONTWRITEBYTECODE=1 python -m unittest",
])
def test_test_runner_options(command):
    assert _is_test_command(command)


@pytest.mark.parametrize("command", [
    "echo 'python3 -B -m unittest'", "pip install pytest", "grep pytest log.txt",
    'python -c "print(123)"', "python3 -c 'import pytest'", "echo 'hello; pytest'",
])
def test_mentions_of_tests_are_not_test_runs(command):
    assert not _is_test_command(command)


async def test_role_classifier_uses_compatible_request(monkeypatch):
    from types import SimpleNamespace

    import rune.agent.litellm_adapter as adapter
    from rune.agent.provenance import classify_roles

    completion = AsyncMock(return_value=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({"source.xlsx": "input", "report.pptx": "output"})))]))
    monkeypatch.setattr(adapter.litellm, "acompletion", completion)
    assert await classify_roles("make report.pptx from source.xlsx", ["source.xlsx", "report.pptx"], "openai/gpt-6-astra", None) == {"source.xlsx": "input", "report.pptx": "output"}
    params = completion.call_args.kwargs
    assert params["model"] == "openai/responses/gpt-6-astra"
    assert params["max_completion_tokens"] == 1200
    assert "max_tokens" not in params


def test_paginated_slides_keep_long_text_and_every_table_cell(tmp_path):
    from pptx import Presentation

    from rune.capabilities.document import DocBlock, DocumentCreateParams
    from rune.capabilities.document_slides import render_slides, verify_slides

    rows = [["주문", "기준일", "팀", "확정금액"]] + [[str(i), "2026-09-07", "영업지원", f"{i},200,000"] for i in range(35)]
    spec = DocumentCreateParams(path="report.pptx", format="pptx", title="9월 확정 매출", blocks=[
        DocBlock(type="heading", text="주의사항"),
        DocBlock(type="bullets", items=["내부 검토용 · 외부 배포 금지 " * 40, "마지막 주의 문구"]),
        DocBlock(type="page_break"), DocBlock(type="table", rows=rows),
    ])
    output = tmp_path / "report.pptx"
    render_slides(output, spec)
    verify_slides(output, spec)
    presentation = Presentation(output)
    tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) > 1
    actual_rows = []
    for table in tables:
        assert [cell.text for cell in table.rows[0].cells] == rows[0]
        actual_rows.extend([[cell.text for cell in row.cells] for row in list(table.rows)[1:]])
    assert actual_rows == rows[1:]
    tables[-1].cell(1, 3).text = "999"
    presentation.save(output)
    with pytest.raises(ValueError, match="changed table cell"):
        verify_slides(output, spec)


def test_wide_table_preserves_row_and_column_identity(tmp_path):
    from rune.capabilities.document import DocBlock, DocumentCreateParams
    from rune.capabilities.document_slides import render_slides, verify_slides
    rows = [[f"column_{i}" for i in range(14)], *[[f"{i}-{j}" for j in range(14)] for i in range(3)]]
    spec = DocumentCreateParams(path="wide.pptx", format="pptx", blocks=[DocBlock(type="table", rows=rows)])
    output = tmp_path / "wide.pptx"
    render_slides(output, spec)
    verify_slides(output, spec)


def test_passing_unittest_updates_completion_freshness():
    from rune.agent.verification_state import VerificationState
    state = VerificationState()
    state.changed()
    assert state.observe_command("python3 -B -m unittest discover -s tests -v", True, "Ran 19 tests in 0.123s\n\nOK")
    assert state.tests_passed_after_edit is True
    state.changed()
    assert state.tests_passed_after_edit is False
    assert not state.observe_command("python3 -B -m unittest", False, "Ran 19 tests\nFAILED (failures=1)")
    assert state.pending


@pytest.mark.parametrize(("command", "expected"), [
    ("python3 - <<'PY'\nprint('hello')\nPY", ""),
    ("python3 -B -m unittest", "unittest"),
    ("uv run ruff check .", "ruff"),
    ("python3 -c 'print(1)'", ""),
    ("echo pytest", ""),
])
def test_suggestions_do_not_learn_shell_fragments(command, expected):
    from rune.utils.shell_command import command_name
    assert command_name(command) == expected


def test_harness_run_checks_update_freshness_at_execution_time(tmp_path):
    from rune.agent.verification_state import VerificationState
    state = VerificationState()
    result = stream(tmp_path, verification_callback=state.observe_command)
    state.changed()
    result._record_mechanical_check("python3 -m pytest", "5 passed", 0)
    assert state.tests_passed_after_edit is True
    state.changed()
    assert state.tests_passed_after_edit is False
    result._record_mechanical_check("python3 -m pytest", "1 failed, 4 passed", 1)
    assert state.pending
    result._record_mechanical_check("python3 -m pytest", "no tests ran", 0)
    assert state.tests_passed_after_edit is False


async def test_concurrent_tool_results_keep_their_own_paths(tmp_path, monkeypatch):
    import asyncio

    from rune.agent.goal_classifier import ClassificationResult
    from rune.agent.loop import NativeAgentLoop
    from rune.types import CapabilityResult

    monkeypatch.setenv("RUNE_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("RUNE_IN_BEST_OF", "1")
    captured = []
    monkeypatch.setattr("rune.agent.loop.build_tool_set", lambda options: captured.append(options) or {})
    loop = NativeAgentLoop()
    await loop._execute_loop("inspect files", "", [], 0, ClassificationResult(goal_type="chat", confidence=1, tier=1), context={"workspace_root": str(tmp_path)})
    options = captured[0]
    first_started, second_finished = asyncio.Event(), asyncio.Event()

    async def first():
        await options.on_tool_start("file_read", {"path": "first.md"})
        first_started.set()
        await second_finished.wait()
        await options.on_tool_end("file_read", CapabilityResult(success=True, output="first"))

    async def second():
        await first_started.wait()
        await options.on_tool_start("file_read", {"path": "second.md"})
        await options.on_tool_end("file_read", CapabilityResult(success=True, output="second"))
        second_finished.set()

    await asyncio.gather(first(), second())
    assert loop._files_read == {"first.md", "second.md"}
