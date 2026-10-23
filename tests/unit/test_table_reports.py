"""Reports must agree with source aggregates, including labels and units."""

import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest

from rune.agent.table_acceptance import TableAcceptance
from rune.agent.table_reports import ReportFacts, compare_facts, evaluate_facts
from rune.capabilities.table_acceptance import TablePlan
from rune.types import CapabilityResult


def plan():
    return TablePlan(applicable=True, requirements=["Save summary.csv and report.docx with every total."],
                     group_by=["team"], aggregates=[{"name": "amount", "column": "amount", "operation": "sum", "unit": "USD"}],
                     grand_total=["TOTAL"], output_names=["summary.csv"],
                     reports=[{"filename": "report.docx", "coverage": "all"}])


def facts(amount="100", unit="USD"):
    return ReportFacts(facts=[{"group": [group], "column": "amount", "value": value,
                              "unit": unit, "quote": f"{group}: {value} {unit}", "label_quotes": [group]}
                             for group, value in [("Engineering", amount), ("Sales", "200"), ("TOTAL", "300")]], complete=True)


EXPECTED = [["engineering", "100"], ["sales", "200"], ["TOTAL", "300"]]


@pytest.mark.parametrize("broad_quote", [False, True])
def test_labels_elsewhere_cannot_ground_a_swapped_amount(broad_quote):
    text = "Engineering: 200 USD\nSales: 100 USD\nTOTAL: 300 USD"
    extracted = facts()
    extracted.facts[0].quote = text if broad_quote else "Sales: 100 USD"
    extracted.facts[1].quote = text if broad_quote else "Engineering: 200 USD"
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == "inconclusive"


def test_unit_elsewhere_does_not_verify_the_claim():
    extracted = facts()
    text = "\n".join(f.quote for f in extracted.facts)
    extracted.facts[0].quote = "Engineering: 100"
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == "inconclusive"


@pytest.mark.parametrize("format", ["table", "heading", "department_total"])
def test_bound_labels_and_units_accept_valid_report_layouts(format):
    extracted = facts()
    for fact in extracted.facts:
        if format == "table":
            fact.quote = f"{fact.group[0]}\t{fact.value} USD"
        elif format == "heading":
            fact.quote = f"{fact.group[0]}\nAmount: {fact.value} USD"
        else:
            fact.quote = f"{fact.group[0]} total: {fact.value} USD"
    text = "\n".join(f.quote for f in extracted.facts)
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == "pass"


def test_currency_heading_must_belong_to_the_numeric_column():
    extracted = facts()
    for fact in extracted.facts:
        fact.quote = f"{fact.group[0]}\t{fact.value}"
    text = "Team\tAmount (USD)\n" + "\n".join(f.quote for f in extracted.facts)
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == "pass"
    text = text.replace("Team\tAmount (USD)", "Team (USD)\tAmount (KRW)")
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == "inconclusive"


@pytest.mark.parametrize("swapped", [False, True])
def test_equal_group_labels_do_not_hide_swapped_aggregate_columns(swapped):
    spec = plan()
    spec.aggregates.append(type(spec.aggregates[0])(name="tax", column="tax", operation="sum"))
    expected = [["engineering", "100", "20"], ["sales", "200", "40"], ["TOTAL", "300", "60"]]
    extracted = facts()
    lines = ["Team\tAmount (USD)\tTax"]
    for fact, row in zip(extracted.facts, expected, strict=True):
        fact.quote = "\t".join([fact.group[0], *reversed(row[1:])]) if swapped else "\t".join([fact.group[0], *row[1:]])
        lines.append(fact.quote)
    # Mentioned coverage still cannot accept a value taken from a different numeric column.
    assert evaluate_facts(spec, expected, "\n".join(lines), extracted, "mentioned").status == ("inconclusive" if swapped else "pass")


def test_currency_identity_keeps_its_literal_evidence():
    spec = plan()
    spec.aggregates[0].unit = "KRW"
    extracted = facts(unit="KRW")
    for fact in extracted.facts:
        fact.unit_quote = "원"
        fact.quote = fact.quote.replace(" KRW", "원")
    text = "\n".join(f.quote for f in extracted.facts)
    assert not compare_facts(spec, EXPECTED, text, extracted, "all")
    assert compare_facts(spec, EXPECTED, text.replace("원", "달러"), extracted, "all")


def test_source_currency_does_not_invent_a_display_requirement():
    spec = plan()
    extracted = facts()
    for fact in extracted.facts:
        fact.quote = fact.quote.replace(" USD", "")
        fact.unit = None
    text = "\n".join(f.quote for f in extracted.facts)
    assert not compare_facts(spec, EXPECTED, text, extracted, "all")
    spec.aggregates[0].unit_required = True
    assert compare_facts(spec, EXPECTED, text, extracted, "all")


@pytest.mark.parametrize("suffix", [".", ", then continue", "; checked"])
def test_sentence_punctuation_is_not_part_of_the_amount(suffix):
    extracted = facts()
    for fact in extracted.facts:
        fact.quote = fact.quote.replace(" USD", suffix)
        fact.unit = None
    assert not compare_facts(plan(), EXPECTED, "\n".join(f.quote for f in extracted.facts), extracted, "all")


@pytest.mark.parametrize("case", ["correct", "amount", "label", "unit", "missing", "invented", "substring", "incomplete"])
def test_claims_are_compared_numerically_after_grounding(case):
    extracted = facts()
    if case == "amount":
        extracted = facts("99")
    elif case == "unit":
        extracted = facts(unit="KRW")
    elif case == "label":
        extracted.facts[0].group = ["Sales"]
        extracted.facts[0].label_quotes = ["Sales"]
        extracted.facts[0].quote = "Sales: 100 USD"
    elif case == "missing":
        extracted.facts.pop()
    elif case == "incomplete":
        extracted.complete = False
    text = "\n".join(f.quote for f in extracted.facts)
    if case == "invented":
        text = text.replace("100", "99")
    elif case == "substring":
        text = text.replace("100", "1000")
        extracted.facts[0].quote = "Engineering: 1000 USD"
    issues = compare_facts(plan(), EXPECTED, text, extracted, "all")
    assert bool(issues) is (case != "correct")


async def test_saved_report_is_rechecked_only_when_bytes_change(tmp_path, monkeypatch):
    from docx import Document

    guard = SimpleNamespace(validate_file_read_path=lambda _: SimpleNamespace(allowed=True))
    monkeypatch.setattr("rune.safety.guardian.get_guardian", lambda: guard)
    source, output, report = tmp_path / "source.csv", tmp_path / "summary.csv", tmp_path / "report.docx"
    source.write_text("team,amount\nengineering,100\nsales,200\n")
    output.write_text("team,amount\nengineering,100\nsales,200\nTOTAL,300\n")
    current = facts("99")

    def save():
        doc = Document()
        for fact in current.facts:
            doc.add_paragraph(fact.quote)
        doc.save(report)

    save()
    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", AsyncMock(return_value=plan()))
    extraction = AsyncMock(side_effect=lambda *args: current.model_dump_json())
    monkeypatch.setattr("rune.agent.requirement_gate._completion", extraction)
    state = TableAcceptance(plan().requirements[0], required=True, workspace=str(tmp_path))
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    state.observe("document_create", {"path": str(report)}, CapabilityResult(success=True))
    assert (await state.verify(contract["id"], str(output), None, 1)).success
    assert "Incorrect amount" in await state.blocker()
    assert state.snapshot()["status"] == "fail"
    assert "document_create" in state.recovery_tools()
    assert await state.blocker()
    assert extraction.await_count == 1
    current = facts()
    save()
    state.observe("document_create", {"path": str(report)}, CapabilityResult(success=True))
    assert state.recovery_tools() is None
    assert await state.blocker() is None
    assert extraction.await_count == 2
    assert state.snapshot()["reports"][0]["status"] == "pass"
    assert await state.blocker() is None
    assert extraction.await_count == 2

    payload = json.loads(extraction.call_args.args[1])
    assert "100" not in json.dumps(payload["SCHEMA"])
    source.write_text("team,amount\nengineering,900\nsales,200\n")
    assert await state.blocker()
    assert extraction.await_count == 2


@pytest.mark.parametrize("case,status", [("correct", "pass"), ("wrong", "fail"), ("partial", "inconclusive"), ("ungrounded", "inconclusive")])
def test_mismatch_requires_a_grounded_contradiction(case, status):
    extracted = facts("99" if case in {"wrong", "ungrounded"} else "100")
    extracted.complete = case != "partial"
    text = "\n".join(f.quote for f in extracted.facts)
    if case == "ungrounded":
        text = text.replace("99", "100")
    assert evaluate_facts(plan(), EXPECTED, text, extracted, "all").status == status


@pytest.mark.parametrize("recovers", [True, False])
@pytest.mark.parametrize("unresolved", ["unavailable", "inconclusive"])
async def test_report_review_retry_is_explicit_bounded_and_never_requests_a_rewrite(tmp_path, monkeypatch, recovers, unresolved):
    from docx import Document

    guard = SimpleNamespace(validate_file_read_path=lambda _: SimpleNamespace(allowed=True))
    monkeypatch.setattr("rune.safety.guardian.get_guardian", lambda: guard)
    source, output, report = tmp_path / "source.csv", tmp_path / "summary.csv", tmp_path / "report.docx"
    source.write_text("team,amount\nengineering,100\nsales,200\n")
    output.write_text("team,amount\nengineering,100\nsales,200\nTOTAL,300\n")
    doc = Document()
    for fact in facts().facts:
        doc.add_paragraph(fact.quote)
    doc.save(report)
    original = report.read_bytes()
    monkeypatch.setattr("rune.agent.table_acceptance.extract_plan", AsyncMock(return_value=plan()))
    ungrounded = facts()
    ungrounded.facts[0].quote = "Sales: 100 USD"
    first = None if unresolved == "unavailable" else ungrounded.model_dump_json()
    extraction = AsyncMock(side_effect=[first, facts().model_dump_json() if recovers else first])
    monkeypatch.setattr("rune.agent.requirement_gate._completion", extraction)
    state = TableAcceptance(plan().requirements[0], required=True, workspace=str(tmp_path))
    contract = (await state.requirements(str(source), None)).metadata["tableContract"]
    await state.verify(contract["id"], str(output), None, 1)
    for _ in range(3):
        assert await state.blocker()
    assert state.snapshot()["status"] == unresolved
    assert state.recovery_tools() == {"table_verify", "task_blocked", "ask_user"}
    assert extraction.await_count == 1
    await state.verify(contract["id"], str(output), None, 1)
    assert (await state.blocker() is None) is recovers
    assert state.snapshot()["status"] == ("pass" if recovers else unresolved)
    if not recovers:
        assert state.recovery_tools() == {"task_blocked", "ask_user"}
        await state.verify(contract["id"], str(output), None, 1)
        assert await state.blocker()
    assert extraction.await_count == 2
    assert report.read_bytes() == original
