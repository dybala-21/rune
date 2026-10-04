"""Local office and approval outcomes with real models and disposable fixtures."""

import asyncio
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock

from docx import Document

from rune.connectors.broker import create_broker
from rune.connectors.models import ConnectorPolicy
from rune.connectors.store import ConnectorStore, authority_key
from rune.types import CapabilityResult
from tests.e2e import test_live_workflows as workflows
from tests.e2e.workflow_harness import TERMINAL, gui_calls
from tests.integration.test_cloud_transport import serving

workflow = workflows.workflow


def test_office_document_preserves_source_and_values(workflow, request):
    source = "출시 실행 계획\n출시일: 2026-10-12\n결제 검증 | 담당: 지민 | 기한: 2026-10-08\n도움말 검수 | 담당: 서준 | 기한: 2026-10-09\n배포 승인 | 담당: 민서 | 기한: 2026-10-11\n"
    (workflow.workspace / "brief.txt").write_text(source)
    run = workflow.submit("brief.txt를 읽고 출시 실행 계획을 rollout.docx로 만들어줘. 출시일과 세 업무의 담당자·기한을 그대로 보존하고 표로 정리해. 새로운 사실은 추가하지 말고 원본도 유지해. 저장한 파일을 다시 읽어서 내용이 맞는지 확인한 후 보고해.")
    workflows.completed(run)
    assert not gui_calls(run)
    assert (workflow.workspace / "brief.txt").read_text() == source
    document = Document(workflow.workspace / "rollout.docx")
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for row in t.rows for c in row.cells])
    assert "2026-10-12" in text
    rows = [[c.text for c in row.cells] for table in document.tables for row in table.rows]
    for task, owner, date in (("결제 검증", "지민", "2026-10-08"), ("도움말 검수", "서준", "2026-10-09"), ("배포 승인", "민서", "2026-10-11")):
        assert any(all(value in " ".join(row) for value in (task, owner, date)) for row in rows), rows
    request.node.live_passed = True


async def test_denied_connector_write_has_no_effect(workflow, monkeypatch, request):
    send = AsyncMock(return_value=CapabilityResult(success=True, output='{"id":"draft-1"}'))
    monkeypatch.setattr("rune.connectors.broker.send_request", send)
    with tempfile.TemporaryDirectory(prefix="rune-denied-e2e-", dir="/private/tmp") as directory:
        root = Path(directory)
        monkeypatch.setenv("RUNE_BROKER_HOME", str(root))
        store = ConnectorStore(root)
        store.put(ConnectorPolicy(name="ledger", origin="https://ledger.example", path_prefix="/v1/drafts", methods=["POST"]), "synthetic-denial-credential")
        async with serving(create_broker(store, authority_key(root, create=True)), uds=root / "broker.sock"):
            run = await asyncio.to_thread(workflow.submit,
                '연결된 ledger API의 /v1/drafts에 POST로 {"title":"E2E draft"}를 보내 임시 기록 하나를 생성해줘. 커넥터 목록에서 연결을 확인해. 내가 승인을 거절하면 다른 도구나 경로로 재시도하지 말고 중단 사실만 알려줘.',
                deny_approvals=True)
    assert workflow.runs[0]["denied_fixture_actions"]
    send.assert_not_awaited()
    assert run["status"] in TERMINAL and not gui_calls(run)
    calls = [c for c in run["toolCalls"] if c["toolName"] == "connector_request"]
    assert len(calls) == 1 and calls[0]["success"] is False
    assert not workflow.runs[0]["approved_fixture_edits"]
    request.node.live_passed = True
