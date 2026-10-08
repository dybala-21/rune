"""Real-model checks for code repair, office output and browser continuity."""

import subprocess
import sys
from decimal import Decimal

from docx import Document
from openpyxl import load_workbook

from tests.e2e.acceptance_fixtures import CHECKS, EXPENSES, INVOICES, MONEY, POLICY
from tests.e2e.acceptance_fixtures import booking_site as booking_site
from tests.e2e.test_live_workflows import completed
from tests.e2e.test_live_workflows import workflow as workflow
from tests.e2e.workflow_harness import gui_calls


def test_invoice_repair_across_modules(workflow, request):
    for name, content in (("money.py", MONEY), ("invoices.py", INVOICES), ("test_invoices.py", CHECKS)):
        (workflow.workspace / name).write_text(content)
    run = workflow.submit(
        "청구서 합계가 틀립니다. money.py와 invoices.py를 검토해서 고쳐줘. "
        "각 항목의 금액은 수량을 곱한 뒤 소수 둘째 자리까지 ROUND_HALF_UP으로 반올림하고, "
        "status가 cancelled인 항목은 합계에서 제외해야 해. 음수 환불도 같은 반올림 규칙을 적용해. "
        "기존 test_invoices.py는 수정하지 말고 수정 전후 테스트를 실행해. 결과와 수정 이유를 간단히 알려줘.",
        editable_files=("money.py", "invoices.py"), allow_new_tests=True,
    )
    completed(run)
    assert not gui_calls(run)
    assert (workflow.workspace / "test_invoices.py").read_text() == CHECKS
    checked = subprocess.run([sys.executable, "-B", "-m", "unittest", "-v", "test_invoices"],
                             cwd=workflow.workspace, capture_output=True, text=True, timeout=15)
    assert checked.returncode == 0, checked.stderr
    probe = subprocess.run([sys.executable, "-B", "-c", '''from decimal import Decimal
from invoices import invoice_total
from money import line_total
assert line_total("2.345", 3) == Decimal("7.04")
assert invoice_total([{"price":"1000","quantity":2,"status":"cancelled"},
                      {"price":"-0.005","quantity":1,"status":"paid"}]) == Decimal("-0.01")
'''], cwd=workflow.workspace, capture_output=True, text=True, timeout=15)
    assert probe.returncode == 0, probe.stderr
    verification = run["trust"]["verification"]
    assert verification["tests_passed_after_edit"] is True
    history = verification["history"]
    assert any(check["write_sequence"] == 0 and check.get("report")
               and check["report"]["failure_events"] == 4 for check in history), verification
    assert any(check["sequence"] > verification["last_write"] and check.get("report")
               and check["report"]["tests_run"] >= 7 and check["report"]["failure_events"] == 0
               and check["status"] == "pass" for check in history), verification
    request.node.live_passed = True


def test_expense_reconciliation_delivers_workbook_and_document(workflow, request):
    for name, content in (("expenses.csv", EXPENSES), ("policy.md", POLICY)):
        (workflow.workspace / name).write_text(content)
    run = workflow.submit(
        "expenses.csv와 policy.md를 읽고 월말 정산 자료를 만들어줘. "
        "reconciliation.xlsx의 Departments 시트에 department,total 두 열로 부서별 정산액을 "
        "이름 오름차순으로 저장해. 마지막 행은 TOTAL로 전체 합계를 넣어. "
        "brief.docx에는 부서별 금액과 전체 합계, 중복·미승인 내역을 어떻게 처리했는지 짧게 정리해. "
        "원본은 유지하고 두 결과 파일을 다시 열어서 수치가 맞는지 검증해줘.",
        editable_files=("reconciliation.xlsx", "brief.docx"),
    )
    completed(run)
    assert not gui_calls(run)
    assert (workflow.workspace / "expenses.csv").read_text() == EXPENSES
    assert (workflow.workspace / "policy.md").read_text() == POLICY
    workbook = load_workbook(workflow.workspace / "reconciliation.xlsx", data_only=True)
    try:
        rows = list(workbook["Departments"].values)
        assert rows[0] == ("department", "total")
        assert [(row[0], Decimal(str(row[1]))) for row in rows[1:]] == [
            ("engineering", Decimal(1000)), ("sales", Decimal(1000)),
            ("support", Decimal(500)), ("TOTAL", Decimal(2500)),
        ], rows
    finally:
        workbook.close()
    doc = Document(workflow.workspace / "brief.docx")
    text = "\n".join([p.text for p in doc.paragraphs] +
                     [" ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
    normalized = text.replace(",", "").casefold()
    assert all(value in normalized for value in ("engineering", "sales", "support", "1000", "500", "2500")), text
    reports = run["trust"]["tableAcceptance"].get("reports", [])
    assert any(report["status"] == "pass" and report["path"].endswith("brief.docx") for report in reports), reports
    request.node.live_passed = True


def test_booking_modal_and_followup_preserve_state(workflow, booking_site, request):
    url, state = booking_site
    first = workflow.submit(
        f"브라우저에서 {url} 연습 화면을 열어. 오전 10시에 6명이 사용할 수 있고 프로젝터가 있는 "
        "회의실을 찾아서 제목 Release review로 임시 예약해줘. 화면에서 검토·저장을 거쳐 한 번만 저장하고 확인해.")
    completed(first)
    assert gui_calls(first)
    draft = {"room": "Cedar", "time": "10:00", "people": "6", "title": "Release review"}
    assert state["saves"] == [draft], state
    second = workflow.submit("열어둔 예약 화면에서 인원만 7명으로 바꿔줘. 나머지 값은 그대로 두고 검토·저장을 한 번 더 해. 같은 화면에서 이어서 진행해.")
    completed(second)
    assert gui_calls(second)
    assert state == {"loads": 1, "saves": [draft, {**draft, "people": "7"}]}, state
    request.node.live_passed = True
