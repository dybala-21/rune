"""Compose expertise and inspect a real office artifact through the web API."""

from tests.e2e.test_live_workflows import completed
from tests.e2e.test_live_workflows import workflow as workflow


def test_two_stage_skills_produce_checked_office_document(workflow, request):
    from docx import Document

    from rune.config import get_config

    get_config().skills.auto_skill = True
    source = workflow.workspace / "release.txt"
    original = "Release 2026-10-12\nPayment check | Mina | 2026-10-08\nHelp review | Sam | 2026-10-09\n"
    source.write_text(original)
    skills = {
        "release-review": ("Read a release source and check owners and deadlines before writing a report.",
                           "Read release.txt. Preserve every name and date exactly. Save review.txt listing the two tasks and append the exact review marker REVIEW-COMPLETE-731."),
        "release-document": ("Create and inspect a Word release report after reviewing a release source.",
                             "Read the completed review.txt. Create release.docx with the release date and a table containing task, owner and deadline. Add the exact report footer RELEASE-REPORT-284. Reopen the saved Word file with document_read."),
    }
    for name, (description, body) in skills.items():
        path = workflow.workspace / ".rune" / "skills" / name / "SKILL.md"
        path.parent.mkdir(parents=True)
        path.write_text(f"---\nname: {name}\ndescription: {description}\n---\n{body}")
    run = workflow.submit("Use the applicable project procedures to review release.txt, then prepare and check the release report as a Word document. Preserve the source facts.")
    completed(run)
    assert "REVIEW-COMPLETE-731" in (workflow.workspace / "review.txt").read_text()
    doc = Document(workflow.workspace / "release.docx")
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert all(value in text for value in ("Mina", "Sam", "2026-10-12", "2026-10-08", "2026-10-09", "RELEASE-REPORT-284"))
    rows = [[c.text for c in row.cells] for table in doc.tables for row in table.rows]
    for task, owner, date in (("Payment check", "Mina", "2026-10-08"), ("Help review", "Sam", "2026-10-09")):
        assert any(all(value in " ".join(row) for value in (task, owner, date)) for row in rows)
    loaded = [call["args"]["name"] for call in run["toolCalls"] if call["toolName"] == "skill_load" and call["success"]]
    assert set(loaded) == set(skills)
    reads = [call for call in run["toolCalls"] if call["toolName"] == "document_read" and call["success"]]
    assert reads and any("Saved-file inspection" in str(call) for call in reads)
    assert not any(call["toolName"].startswith(("browser_", "desktop_")) for call in run["toolCalls"])
    assert source.read_text() == original
    request.node.live_passed = True
