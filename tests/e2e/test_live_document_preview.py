"""The model must receive actual rendered pixels to inspect a title color."""

import hashlib
import json
import shutil

import pytest

from tests.e2e.test_live_workflows import completed
from tests.e2e.test_live_workflows import workflow as workflow


@pytest.mark.skipif(not shutil.which("soffice") or not shutil.which("pdftoppm"), reason="Headless renderers not installed")
def test_rendered_page_reaches_model(workflow, request):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_paragraph("First page")
    doc.add_page_break()
    run = doc.add_paragraph().add_run("RELEASE")
    run.font.size = Pt(60)
    run.font.color.rgb = RGBColor(255, 0, 0)
    path = workflow.workspace / "preview.docx"
    doc.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    run = workflow.submit("Use document_preview to inspect page 2 of preview.docx. Report the rendered document's page count and the large title's color. Reply only as JSON with page_count and title_color (an English color name). Do not modify the document.")
    completed(run)
    answer = (run.get("answer") or run["text"]).strip()
    if answer.startswith("```"):
        answer = "\n".join(answer.splitlines()[1:-1])
    result = json.loads(answer)
    assert result["page_count"] == 2 and result["title_color"].lower() == "red"
    assert any(call["toolName"] == "document_preview" and call["success"] for call in run["toolCalls"])
    assert not any(call["toolName"].startswith(("browser_", "desktop_")) for call in run["toolCalls"])
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    request.node.live_passed = True
