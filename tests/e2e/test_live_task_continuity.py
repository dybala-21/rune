"""Continue coding and office work in a new chat with changed requirements."""

import subprocess
from uuid import uuid4

import pytest
from docx import Document

from rune.config import get_config
from rune.memory.task_store import TaskStore
from tests.e2e.test_live_workflows import completed
from tests.e2e.test_live_workflows import workflow as workflow
from tests.e2e.workflow_harness import gui_calls


def _limits():
    get_config().llm.cost_budget_usd = .30
    get_config().llm.model_request_limit = 14


def _saved_task(workflow):
    tasks = [task for task in TaskStore().list('local', str(workflow.workspace.resolve()))
             if task.get('kind') == 'conversation']
    assert len(tasks) == 1, tasks
    return tasks[0]['id']


def _check_invoice_precision(workspace):
    checked = subprocess.run(['node', '--input-type=module', '-e', """
import assert from 'node:assert/strict';
import {total} from './invoice.mjs';
const cases = [
  [[.2, .4], .6], [[.001, .002], .003], [[1.234], 1.234],
  [[-.001, .004], .003], [[.002, .001], .003],
];
for (let scale = 1; scale <= 16; scale++) {
  for (const sign of [1, -1]) {
    cases.push([[Number(`${sign}e-${scale}`), Number(`2e-${scale}`)], Number(`${sign + 2}e-${scale}`)]);
  }
}
cases.push([[1.23456789012345], 1.23456789012345], [[1e15, .001, -1e15], .001]);
for (const [amounts, expected] of cases) {
  const rows = amounts.map(amount => ({status: 'paid', amount}));
  rows.push({status: 'refunded', amount: 99});
  assert.equal(total(rows), expected, `Precision changed for ${amounts}`);
}
"""], cwd=workspace, capture_output=True, text=True, timeout=20)
    assert checked.returncode == 0, checked.stdout + checked.stderr


@pytest.mark.parametrize('workflow', [{'max_iterations': 14, 'timeout_seconds': 180,
                                       'cost_budget_usd': .30}], indirect=True)
def test_javascript_fix_keeps_the_task_across_chats(workflow, request):
    _limits()
    source = 'export function total(rows) { return rows.reduce((sum, row) => sum + row.amount, 0); }\n'
    tests = '''import test from 'node:test';
import assert from 'node:assert/strict';
import { total } from './invoice.mjs';
test('paid invoices', () => assert.equal(total([{status:'paid', amount:.2}, {status:'refunded', amount:99}, {status:'paid', amount:.4}]), .6));
test('no invoices', () => assert.equal(total([]), 0));
test('integer total', () => assert.equal(total([{status:'paid', amount:20}, {status:'paid', amount:5.5}]), 25.5));
'''
    path = workflow.workspace / 'invoice.mjs'
    path.write_text(source)
    (workflow.workspace / 'invoice.test.mjs').write_text(tests)
    first = workflow.submit('invoice.mjs의 청구 합계 버그를 맡아줘. paid 행만 합산하고 공개 함수 이름과 인자를 유지해야 해. '
                            '지금은 파일을 읽고 수정할 사항만 파악해. 코드와 테스트 파일은 아직 변경하지 마.')
    completed(first)
    assert path.read_text() == source
    task_id = _saved_task(workflow)
    TaskStore().create('local', str(workflow.workspace.resolve()), [
        {'id': uuid4().hex[:12], 'goal': f'Prepare unrelated presentation {index}',
         'status': 'pending', 'kind': 'manual'} for index in range(8)
    ])
    workflow.session = uuid4().hex
    second = workflow.submit('앞서 맡긴 청구 합계 수정을 이제 진행해. 추가로 소수 금액도 정확히 계산돼야 해. '
                             '기존 테스트는 변경하지 말고 node --test로 실행해서 결과를 확인해.',
                             editable_files=('invoice.mjs',))
    completed(second)
    assert (workflow.workspace / 'invoice.test.mjs').read_text() == tests
    checked = subprocess.run(['node', '--test', 'invoice.test.mjs'], cwd=workflow.workspace,
                             capture_output=True, text=True, timeout=20)
    assert checked.returncode == 0, checked.stdout + checked.stderr
    _check_invoice_precision(workflow.workspace)
    assert _saved_task(workflow) == task_id
    assert not gui_calls(first) and not gui_calls(second)
    request.node.live_passed = True


@pytest.mark.parametrize('workflow', [{'max_iterations': 14, 'timeout_seconds': 180,
                                       'cost_budget_usd': .30}], indirect=True)
def test_office_document_uses_revised_source_in_a_new_chat(workflow, request):
    _limits()
    path = workflow.workspace / 'launch.txt'
    original = ('출시일: 2026-11-10\n결제 검증 | 지민 | 2026-11-06 | 진행\n'
                '도움말 검수 | 서준 | 2026-11-07 | 진행\n이전 행사 | 민서 | 2026-11-08 | 취소\n')
    path.write_text(original)
    first = workflow.submit('launch.txt를 바탕으로 출시 실행 계획 Word 문서를 만드는 일을 맡아줘. '
                            '출시일을 표시하고 업무·담당자·기한을 표로 정리해야 해. 원본은 보존해. '
                            '지금은 자료만 읽고 아직 문서를 만들지는 마.')
    completed(first)
    assert not (workflow.workspace / 'launch.docx').exists()
    task_id = _saved_task(workflow)
    updated = original.replace('서준 | 2026-11-07', '민서 | 2026-11-09')
    path.write_text(updated)
    TaskStore().create('local', str(workflow.workspace.resolve()), [
        {'id': uuid4().hex[:12], 'goal': f'Prepare unrelated presentation {index}',
         'status': 'pending', 'kind': 'manual'} for index in range(8)
    ])
    workflow.session = uuid4().hex
    second = workflow.submit('앞서 맡긴 출시 계획 문서를 이제 launch.docx로 만들어줘. 원본 자료를 갱신했으니 다시 읽고 '
                             '최신 담당자와 기한을 반영해. 추가 조건은 취소된 업무를 표에서 제외하는 거야. '
                             '저장한 문서를 읽어 내용이 맞는지 확인해.')
    completed(second)
    assert path.read_text() == updated
    document = Document(workflow.workspace / 'launch.docx')
    rows = [' '.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    text = '\n'.join([paragraph.text for paragraph in document.paragraphs] + rows)
    assert '2026-11-10' in text
    assert any(all(value in row for value in ('결제 검증', '지민', '2026-11-06')) for row in rows)
    assert any(all(value in row for value in ('도움말 검수', '민서', '2026-11-09')) for row in rows)
    assert all('이전 행사' not in row for row in rows)
    assert _saved_task(workflow) == task_id
    assert not gui_calls(first) and not gui_calls(second)
    request.node.live_passed = True
