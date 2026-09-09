"""Paginate document blocks into editable slides with bounded text and table rows."""

from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from rune.capabilities.document import DocumentCreateParams


def _width(text: str) -> float:
    return sum(
        0 if unicodedata.combining(c) else
        1.0 if unicodedata.east_asian_width(c) in {"W", "F"} else
        0.34 if c.isspace() else 0.68
        for c in text
    )


def _lines(text: str, width_pt: float, size: float) -> list[str]:
    capacity = width_pt / size * 0.90
    lines: list[str] = []
    for line in text.split("\n"):
        current = ""
        for word in line.split():
            if _width(word) > capacity:
                raise ValueError("Slide text contains a word too wide to display; split the content into shorter fields")
            candidate = f"{current} {word}" if current else word
            if current and _width(candidate) > capacity:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    return lines


def render_slides(path: Path, params: DocumentCreateParams) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.core_properties.identifier = "rune:document-slides:1"
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    left, width, top, bottom = 40.0, 880.0, 112.0, 498.0
    family = params.font_family or "Arial"

    def format_frame(frame: Any, size: float) -> None:
        frame.margin_left = frame.margin_right = Pt(7)
        frame.margin_top = frame.margin_bottom = Pt(5)
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for para in frame.paragraphs:
            para.font.name = family
            para.font.size = Pt(size)
            para.font.color.rgb = RGBColor.from_string("17212F")  # type: ignore[no-untyped-call]
            para.line_spacing = Pt(size * 1.4)
            para.space_before = para.space_after = Pt(0)

    def new_slide(title: str, title_name: str = "rune-continuation") -> Any:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if title:
            lines = _lines(title, width - 14, 30)
            if len(lines) > 2:
                raise ValueError("Slide heading exceeds two lines; use a shorter heading")
            box = slide.shapes.add_textbox(Pt(left), Pt(16), Pt(width), Pt(94))
            box.name = title_name
            box.text_frame.text = "\v".join(lines)
            format_frame(box.text_frame, 30)
            box.text_frame.paragraphs[0].font.bold = True
        return slide

    if params.title:
        cover = new_slide("")
        lines = _lines(params.title, width - 14, 42)
        if len(lines) > 6:
            raise ValueError("Presentation title is too long for the cover")
        box = cover.shapes.add_textbox(Pt(left), Pt(110), Pt(width), Pt(390))
        box.name = "rune-title"
        box.text_frame.text = "\v".join(lines)
        format_frame(box.text_frame, 42)

    slide = None
    heading = params.title
    y = top
    for block_index, block in enumerate(params.blocks):
        if block.type == "heading":
            heading = block.text
            slide = new_slide(heading, f"rune-block-{block_index}")
            y = top
        elif block.type == "page_break":
            slide = None
            y = top
        elif block.type in {"paragraph", "bullets"}:
            texts = block.items if block.type == "bullets" else [block.text]
            for item_index, text in enumerate(texts):
                lines = _lines(text, width - 34, 20)
                while lines:
                    if slide is None or bottom - y < 48:
                        slide = new_slide(heading)
                        y = top
                    count = max(1, math.floor((bottom - y - 18) / 28))
                    chunk, lines = lines[:count], lines[count:]
                    height = len(chunk) * 28 + 18
                    box = slide.shapes.add_textbox(Pt(left), Pt(y), Pt(width), Pt(height))
                    box.name = f"rune-block-{block_index}-item-{item_index}"
                    prefix = "• " if block.type == "bullets" else ""
                    box.text_frame.text = prefix + "\v".join(chunk)
                    format_frame(box.text_frame, 20)
                    y += height
        elif block.type == "table" and block.rows:
            cols = max(map(len, block.rows))
            if not cols:
                continue
            rows = [[str(row[i]) if i < len(row) else "" for i in range(cols)] for row in block.rows]
            size = 17.0
            minima = [max(65.0, max((_width(word) for row in rows for word in row[c].split()), default=0) * size / 0.9 + 20) for c in range(cols)]
            # Wide tables continue by columns, keeping the first column as the row key.
            groups: list[list[int]] = []
            group: list[int] = []
            used = 0.0
            for c, minimum in enumerate(minima):
                if minimum > width or (c and minima[0] + minimum > width):
                    raise ValueError("Table cell is too wide for a slide; shorten the field or use a separate detail document")
                if group and used + minimum > width:
                    groups.append(group)
                    group, used = [0], minima[0]
                group.append(c)
                used += minimum
            if group:
                groups.append(group)
            for columns in groups:
                spare = (width - sum(minima[c] for c in columns)) / len(columns)
                widths = [minima[c] + spare for c in columns]
                wrapped = [[_lines(row[c], col_width - 14, size) for c, col_width in zip(columns, widths, strict=True)] for row in rows]
                heights = [max(map(len, row)) * size * 1.4 + 14 for row in wrapped]
                if heights[0] > bottom - top:
                    raise ValueError("Table header is too tall for a slide")
                reuse_slide = (columns == groups[0] and slide is not None
                               and heights[0] + (heights[1] if len(heights) > 1 else 0) <= bottom - y)
                capacity = bottom - y if reuse_slide else bottom - top
                pages: list[list[int]] = []
                page, used = [0], heights[0]
                for index in range(1, len(rows)):
                    if heights[0] + heights[index] > bottom - top:
                        raise ValueError("Table row is too tall for a slide; split the row into smaller records")
                    if used + heights[index] > capacity:
                        pages.append(page)
                        page, used = [0], heights[0]
                        capacity = bottom - top
                    page.append(index)
                    used += heights[index]
                pages.append(page)
                for page_index, indices in enumerate(pages):
                    reuse = reuse_slide and page_index == 0
                    table_slide = slide if reuse else new_slide(heading)
                    assert table_slide is not None
                    table_top = y if reuse else top
                    shape = table_slide.shapes.add_table(len(indices), len(columns), Pt(left), Pt(table_top), Pt(width), Pt(sum(heights[i] for i in indices)))
                    shape.name = f"rune-table-{block_index}-cols-{','.join(map(str, columns))}-rows-{','.join(map(str, indices))}"
                    table = shape.table
                    for c, col_width in enumerate(widths):
                        table.columns[c].width = Pt(col_width)
                    for r, index in enumerate(indices):
                        table.rows[r].height = Pt(heights[index])
                        for c, content in enumerate(wrapped[index]):
                            cell = table.cell(r, c)
                            cell.text = "\v".join(content)
                            format_frame(cell.text_frame, size)
                            if index == 0:
                                cell.fill.solid()
                                cell.fill.fore_color.rgb = RGBColor.from_string("E8EEF5")  # type: ignore[no-untyped-call]
                                cell.text_frame.paragraphs[0].font.bold = True
            slide = None
            y = top
    if not prs.slides:
        new_slide("")
    prs.save(str(path))


def verify_slides(path: Path, params: DocumentCreateParams) -> bool:
    """Read persisted blocks in document order, including paginated tables."""
    from pptx import Presentation

    texts: dict[str, list[str]] = {}
    cells: dict[tuple[int, int, int], str] = {}
    prs = Presentation(str(path))
    if prs.core_properties.identifier != "rune:document-slides:1":
        return False
    for slide in prs.slides:
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > prs.slide_width
                    or shape.top + shape.height > prs.slide_height):
                raise ValueError("Saved slide has a shape outside the page")
            if shape.has_text_frame:
                texts.setdefault(shape.name, []).append(shape.text)
            match = re.fullmatch(r"rune-table-(\d+)-cols-([\d,]+)-rows-([\d,]+)", shape.name)
            if match and shape.has_table:
                block_index = int(match[1])
                columns = list(map(int, match[2].split(",")))
                rows = list(map(int, match[3].split(",")))
                for r, original_row in enumerate(rows):
                    for c, original_col in enumerate(columns):
                        key = block_index, original_row, original_col
                        value = " ".join(shape.table.cell(r, c).text.split())
                        if key in cells and cells[key] != value:
                            raise ValueError("Repeated table cells differ between slides")
                        cells[key] = value

    def check_text(name: str, expected: str, *, bullet: bool = False) -> None:
        chunks = texts.get(name, [])
        if bullet:
            chunks = [chunk.removeprefix("• ") for chunk in chunks]
        if " ".join(" ".join(chunks).split()) != " ".join(expected.split()):
            raise ValueError(f"Saved pptx lost supplied text: {expected[:80]!r}")

    if params.title:
        check_text("rune-title", params.title)
    for i, block in enumerate(params.blocks):
        if block.type == "heading":
            check_text(f"rune-block-{i}", block.text)
        elif block.type in {"paragraph", "bullets"}:
            for j, text in enumerate(block.items if block.type == "bullets" else [block.text]):
                check_text(f"rune-block-{i}-item-{j}", text, bullet=block.type == "bullets")
        elif block.type == "table":
            for r, row in enumerate(block.rows):
                for c, expected_cell in enumerate(row):
                    if cells.get((i, r, c)) != " ".join(str(expected_cell).split()):
                        raise ValueError(f"Saved pptx changed table cell {r + 1},{c + 1}")
    return True
