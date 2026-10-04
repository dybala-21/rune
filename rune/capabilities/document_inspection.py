"""Read saved content and measurable structure without claiming visual quality."""

from __future__ import annotations

import hashlib
import io
import os
import zipfile
from pathlib import Path

from rune.agent.isolation import enforce
from rune.safety.guardian import get_guardian

MAX_BYTES = 10_000_000
MAX_ITEMS = 50_000


def read_snapshot(path: str | Path) -> tuple[Path, bytes]:
    target = Path(path).expanduser().resolve()
    decision = get_guardian().validate_file_read_path(str(target))
    if not decision.allowed:
        raise ValueError(decision.reason)
    if reason := enforce(str(target)):
        raise ValueError(reason)
    if not target.is_file():
        raise ValueError(f"File not found: {target}")
    with target.open("rb") as stream:
        before = os.fstat(stream.fileno())
        data = stream.read(MAX_BYTES + 1)
        after = os.fstat(stream.fileno())
    if (before.st_size, before.st_mtime_ns, before.st_ctime_ns) != (after.st_size, after.st_mtime_ns, after.st_ctime_ns):
        raise ValueError("File changed while it was being inspected")
    if len(data) > MAX_BYTES:
        raise ValueError("Document inspection is limited to 10 MB")
    return target, data


def inspect_document(path: str | Path, max_chars: int = 20_000) -> dict:
    target, data = read_snapshot(path)
    result = inspect_bytes(data, target.suffix.lower().lstrip("."), max_chars)
    if target.suffix.lower() in {".csv", ".tsv"}:
        from rune.capabilities.table_profile import profile_table

        evidence: dict = {}
        result["table_profile_text"] = profile_table(data.decode("utf-8-sig"), target.name, evidence=evidence)
        result["table_profile"] = evidence
    return result


def inspect_bytes(data: bytes, fmt: str, max_chars: int = 20_000) -> dict:
    """All facts and the hash describe the same byte snapshot."""
    facts: dict = {"format": fmt, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}
    warnings: list[str] = []
    lines: list[str] = []
    size = items = 0
    truncated = False

    def add(value: str) -> None:
        nonlocal size, items, truncated
        items += 1
        if items > MAX_ITEMS:
            raise ValueError("Document has too many items to inspect")
        remaining = max(0, max_chars - size)
        if remaining:
            lines.append(value[:remaining])
        size += len(value) + 1
        truncated |= size > max_chars

    stream = io.BytesIO(data)
    if fmt in {"docx", "pptx", "xlsx"}:
        with zipfile.ZipFile(stream) as archive:
            members = archive.infolist()
            if len(members) > 10_000 or sum(m.file_size for m in members) > 50_000_000:
                raise ValueError("Expanded document exceeds the inspection limit")
        stream.seek(0)
    if fmt == "docx":
        from docx import Document
        from docx.text.paragraph import Paragraph

        doc = Document(stream)
        facts.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables), page_count=None)
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                add(block.text)
            else:
                for row in block.rows:
                    add("\t".join(cell.text for cell in row.cells))
        warnings.append("Word pagination, headers, text boxes and visual layout are not verified by body text extraction.")
    elif fmt == "pptx":
        from pptx import Presentation

        prs = Presentation(stream)
        outside = []
        facts["slides"] = len(prs.slides)
        for i, slide in enumerate(prs.slides, 1):
            add(f"# Slide {i}")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text_frame.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        add("\t".join(c.text for c in row.cells))
                if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width
                        or shape.top + shape.height > prs.slide_height):
                    outside.append({"slide": i, "shape": shape.shape_id})
        facts["shapes_outside_slide"] = outside
        warnings.append("Shape bounds do not establish text fit, overlap or visual quality; groups and images need rendered inspection.")
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(stream, read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas = missing = 0
        facts["sheets"] = []
        try:
            for ws in wb.worksheets:
                if (ws.max_row or 0) * (ws.max_column or 0) > MAX_ITEMS:
                    raise ValueError("Worksheet exceeds the inspection cell limit")
                facts["sheets"].append(ws.title)
                add(f"# Sheet: {ws.title}")
                for row, values in zip(ws.iter_rows(), cached[ws.title].iter_rows(), strict=True):
                    cells = []
                    for cell, value in zip(row, values, strict=True):
                        items += 1
                        if items > MAX_ITEMS:
                            raise ValueError("Workbook exceeds the inspection cell limit")
                        if cell.data_type == "f":
                            formulas += 1
                            missing += value.value is None
                            cells.append(f"{cell.value} [cached: {value.value!r}]")
                        else:
                            cells.append("" if cell.value is None else str(cell.value))
                    add("\t".join(cells))
        finally:
            wb.close()
            cached.close()
        facts.update(formulas=formulas, formulas_without_cached_values=missing)
        if formulas:
            warnings.append("Formulas were not recalculated. Cached values may be missing or stale; they are not verified calculations.")
    elif fmt == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(stream)
        if len(reader.pages) > 500:
            raise ValueError("PDF exceeds the 500-page inspection limit")
        facts["page_count"] = len(reader.pages)
        empty = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            add(text)
            if not text.strip():
                empty.append(i)
        facts["pages_without_text"] = empty
        warnings.append("PDF page count is measured; text extraction does not verify clipping, images or visual layout.")
    else:
        add(data.decode("utf-8-sig"))
    if truncated:
        warnings.append("Content is truncated; omitted text is not evidence.")
    return {"text": "\n".join(lines)[:max_chars], "facts": facts,
            "truncated": truncated, "warnings": warnings, "visual_verified": False}
