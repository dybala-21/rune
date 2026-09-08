"""Render and check an office bundle before publishing its version.

Readers resolve current.json once to get the files for one version. Concurrent
publishes compare the version observed before rendering. Checks cover saved content and metric
consistency; they do not validate prose accuracy, formulas, or visual layout.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import tempfile
import uuid
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from rune.agent.isolation import enforce
from rune.capabilities.bundle_data import (
    BundleMetric,
    RowFilter,
    resolve_references,
)
from rune.capabilities.bundle_revision import (
    BundleError,
    BundleSnapshot,
    changes_between,
    check_controls,
    digest,
    publish,
    read_snapshot,
)
from rune.capabilities.document import (
    _READERS,
    DocBlock,
    DocSheet,
    DocumentCreateParams,
)
from rune.capabilities.document import (
    _RENDERERS as _RENDERERS,
)
from rune.capabilities.registry import CapabilityRegistry
from rune.capabilities.types import CapabilityDefinition
from rune.safety.guardian import get_guardian
from rune.types import CapabilityResult, Domain, RiskLevel
from rune.utils.logger import get_logger

log = get_logger(__name__)
BundleFormat = Literal["xlsx", "docx", "pptx", "pdf"]


class BundleDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")
    filename: str = Field(description="Simple filename including format extension, no directories")
    format: BundleFormat
    title: str = ""
    font_family: str = ""
    blocks: list[DocBlock] = Field(default_factory=list)
    sheets: list[DocSheet] = Field(default_factory=list)


class DocumentBundleParams(BaseModel):
    model_config = ConfigDict(extra="forbid")
    directory: str = Field(min_length=1, description="Bundle output folder, with current.json and versions/")
    source_path: str = Field(min_length=1, description="CSV or values-only XLSX source, first row contains headers")
    sheet: str | None = Field(default=None, description="Explicit source sheet if workbook has multiple tabs")
    filters: list[RowFilter] = Field(default_factory=list, description="All filters apply together (AND)")
    metrics: list[BundleMetric] = Field(min_length=1, max_length=100)
    documents: list[BundleDocument] = Field(
        min_length=1, max_length=12,
        description="Use {{metric_id}} references in prose/table cells; resend spec to update all files",
    )


def _hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json(data: object) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False)


def _verify_file(path: Path, params: DocumentCreateParams) -> None:
    """Compare saved cells or extracted text with the requested content."""
    if params.format == "pptx":
        from rune.capabilities.document_slides import verify_slides

        if verify_slides(path, params):
            return
    if params.format == "xlsx":
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            if wb.sheetnames != [s.name for s in params.sheets]:
                raise ValueError("Saved sheet names differ from the requested names")
            for spec in params.sheets:
                actual = list(wb[spec.name].values)
                # Excel reads empty strings as None and pads ragged rows.
                width = max(len(r) for r in spec.rows)
                expected = [tuple(None if v == "" else v for v in row) +
                            (None,) * (width - len(row)) for row in spec.rows]
                if actual != expected:
                    raise ValueError(f"Saved values differ in sheet {spec.name}")
        finally:
            wb.close()
        return
    text = " ".join(_READERS[params.format][0](path).split())
    fragments = [params.title]
    for block in params.blocks:
        if block.type in ("heading", "paragraph"):
            fragments.append(block.text)
        elif block.type == "bullets":
            fragments.extend(block.items)
        elif block.type == "table":
            fragments.extend(str(c) for row in block.rows for c in row)
    cursor = 0
    for fragment in fragments:
        normalized = " ".join(fragment.split())
        position = text.find(normalized, cursor)
        if position < 0:
            raise ValueError(f"Saved {params.format} lost supplied text: {fragment[:80]!r}")
        cursor = position + len(normalized)


def _prepare_documents(
    params: DocumentBundleParams, values: dict[str, int | float], display: dict[str, str],
) -> list[DocumentCreateParams]:
    prepared = []
    names: set[str] = set()
    for doc in params.documents:
        name = doc.filename
        if (not name or name.startswith(".") or "/" in name or "\\" in name
                or Path(name).suffix != f".{doc.format}" or name.casefold() in names):
            raise ValueError(f"Invalid or duplicate document filename: {name}")
        names.add(name.casefold())
        spec = doc.model_dump(exclude={"filename", "format"})
        resolved = resolve_references(spec, values, display)
        # Keep prose fields as text when a reference resolves to a number.
        resolved["title"] = str(resolved["title"])
        for block in resolved["blocks"]:
            block["text"] = str(block["text"])
            block["items"] = [str(v) for v in block["items"]]
        if doc.format == "xlsx":
            if not resolved["sheets"] or any(not s["rows"] for s in resolved["sheets"]):
                raise ValueError("Bundle XLSX requires explicit nonempty sheets")
            for sheet in resolved["sheets"]:
                for row in sheet["rows"]:
                    if any(isinstance(v, str) and v.startswith("=") for v in row):
                        raise ValueError("Bundle output cells must be values, not unevaluated formulas")
        prepared.append(DocumentCreateParams(path=name, format=doc.format, **resolved))
    return prepared


class BundleInspectParams(BaseModel):
    model_config = ConfigDict(extra="forbid")
    directory: str = Field(min_length=1)
    revision: str | None = Field(default=None, pattern=r"^[0-9a-f]{32}$")


class BundleDocumentPatch(BaseModel):
    model_config = ConfigDict(extra="forbid")
    filename: str
    title: str | None = None
    font_family: str | None = None
    blocks: list[DocBlock] | None = None
    sheets: list[DocSheet] | None = None


class BundleChanges(BaseModel):
    model_config = ConfigDict(extra="forbid")
    filters: list[RowFilter] | None = None
    metrics: list[BundleMetric] | None = None
    documents: list[BundleDocumentPatch] | None = None
    source_path: str | None = None
    sheet: str | None = None


class BundleUpdateParams(BaseModel):
    model_config = ConfigDict(extra="forbid")
    directory: str = Field(min_length=1)
    base_revision: str = Field(pattern=r"^[0-9a-f]{32}$")
    expected_source_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    changes: BundleChanges


def _updates_enabled() -> bool:
    return os.environ.get("RUNE_BUNDLE_UPDATE_ENABLED", "").lower() in {"1", "true", "yes", "on"}


def _authorize(root: Path, source: Path | None = None, *, write: bool = True) -> None:
    guardian = get_guardian()
    check = guardian.validate_file_path(str(root)) if write else guardian.validate_file_read_path(str(root))
    if not check.allowed:
        raise BundleError("permission_denied", check.reason)
    if write and (error := enforce(str(root))):
        raise BundleError("permission_denied", error)
    if source is not None:
        check = guardian.validate_file_read_path(str(source))
        if not check.allowed:
            raise BundleError("permission_denied", check.reason)
        if root == source or root in source.parents:
            raise BundleError("invalid_source", "Keep the source outside the bundle output directory")


def _source_bytes(source: Path) -> bytes:
    with source.open("rb") as stream:
        data = stream.read(10_000_001)
    if len(data) > 10_000_000:
        raise BundleError("source_too_large", "Source exceeds 10 MB")
    return data


def _failure(exc: Exception) -> CapabilityResult:
    log.warning("document_bundle_failed", error=str(exc))
    return CapabilityResult(success=False, error=f"Bundle was not published: {exc}", metadata={
        "code": getattr(exc, "code", "bundle_failed"),
        "current_revision": getattr(exc, "current_revision", None),
    })


async def _render_staged(payload: dict[str, Any]) -> dict[str, Any]:
    from rune.capabilities.document_worker import serve
    from rune.utils.process_worker import ProcessWorker

    worker = ProcessWorker(serve)
    try:
        result = await asyncio.to_thread(worker.request, payload, timeout=120)
        if not result.get("success"):
            raise BundleError("render_failed", result.get("error", "Document worker failed"))
        data = result.get("data")
        if not isinstance(data, dict):
            raise BundleError("invalid_render", "Document worker returned invalid data")
        return data
    finally:
        # Close before TemporaryDirectory cleanup, including after cancellation.
        worker.close()


async def _create_version(params: DocumentBundleParams, base: BundleSnapshot | None,
                          expected_source: str | None = None) -> CapabilityResult:
    root = Path(params.directory).expanduser().resolve()
    source = Path(params.source_path).expanduser().resolve()
    _authorize(root, source)
    check_controls(root)
    data = _source_bytes(source)
    source_hash = _hash(data)
    if expected_source is not None and source_hash != expected_source:
        raise BundleError("source_changed", "Source differs from the inspected snapshot")
    spec = params.model_copy(update={"directory": str(root), "source_path": str(source)})
    (root / "versions").mkdir(parents=True, exist_ok=True)
    revision = uuid.uuid4().hex
    final = root / "versions" / revision
    with tempfile.TemporaryDirectory(prefix=".stage-", dir=root) as temp:
        stage = Path(temp)
        (stage / "version").mkdir()
        (stage / "source.snapshot").write_bytes(data)
        rendered = await _render_staged({"stage": str(stage), "spec": spec.model_dump()})
        if _hash(_source_bytes(source)) != source_hash:
            raise BundleError("source_changed", "Source changed during rendering; inspect before retrying")
        artifacts = []
        expected_names = {d.filename for d in spec.documents}
        if {a["filename"] for a in rendered["artifacts"]} != expected_names:
            raise BundleError("invalid_render", "Rendered file list differs from the specification")
        for artifact in rendered["artifacts"]:
            name = artifact["filename"]
            path = stage / "version" / name
            if Path(name).name != name or path.is_symlink() or digest(path) != artifact["sha256"]:
                raise BundleError("invalid_render", "Staged artifact changed after verification")
            artifacts.append({k: v for k, v in artifact.items() if k != "filename"} | {"path": str(final / name)})
        verification = {"status": "pass", "checks": ["native_readback", "metric_resolution"],
                        "visual_review": "not_performed", "task_acceptance": "not_performed"}
        receipt = {"kind": "document_bundle", "revision": revision, "source_sha256": source_hash,
                   "artifacts": [{"path": a["path"], "sha256": a["sha256"]} for a in artifacts],
                   "checks": {"native_content": "pass", "source_metrics": "pass",
                              "visual_layout": "not_performed", "task_acceptance": "not_performed"}}
        manifest = {
            "schema_version": 2, "revision": revision, "parent_revision": base.revision if base else None,
            "source": {"path": str(source), "sha256": source_hash, "sheet": spec.sheet},
            **{k: rendered[k] for k in ("rows_total", "rows_selected", "metrics", "display_values")},
            "spec": spec.model_dump(), "artifacts": artifacts, "verification": verification,
            "receipt": receipt,
        }
        manifest["changes"] = changes_between(base.manifest if base else {}, manifest)
        publish(root, stage, manifest, base)
    paths = [a["path"] for a in artifacts]
    return CapabilityResult(success=True, output=_json({
        "revision": revision, "parent_revision": manifest["parent_revision"],
        "metrics": rendered["metrics"], "files": paths, "manifest": str(final / "manifest.json"),
        "verification": verification, "changes": manifest["changes"], "receipt": receipt,
    }), metadata={"paths": paths + [str(root / "current.json")], "manifest": str(final / "manifest.json"),
                  "revision": revision, "verified": True, "metrics": rendered["metrics"], "receipt": receipt})


async def document_bundle(params: DocumentBundleParams) -> CapabilityResult:
    try:
        root = Path(params.directory).expanduser().resolve()
        _authorize(root)
        return await _create_version(params, read_snapshot(root))
    except Exception as exc:
        return _failure(exc)


async def document_bundle_inspect(params: BundleInspectParams) -> CapabilityResult:
    try:
        root = Path(params.directory).expanduser().resolve()
        _authorize(root, write=False)
        snapshot = read_snapshot(root, params.revision, strict=False)
        if snapshot is None:
            raise BundleError("not_found", "No published bundle exists")
        source = Path(snapshot.manifest["source"]["path"]).expanduser().resolve()
        _authorize(root, source, write=False)
        current_source = _hash(_source_bytes(source)) if source.is_file() else None
        return CapabilityResult(success=True, output=_json({
            "revision": snapshot.revision, "manifest_sha256": snapshot.manifest_sha256,
            "spec": snapshot.manifest["spec"], "metrics": snapshot.manifest["metrics"],
            "source_sha256": snapshot.manifest["source"]["sha256"],
            "current_source_sha256": current_source,
            "source_stale": current_source != snapshot.manifest["source"]["sha256"],
            "artifacts": snapshot.manifest["artifacts"],
            "changed_artifacts": snapshot.changed_artifacts,
            "verification": snapshot.manifest.get("verification", {}),
        }), metadata={"path": str(root / "current.json"), "revision": snapshot.revision})
    except Exception as exc:
        return _failure(exc)


async def document_bundle_update(params: BundleUpdateParams) -> CapabilityResult:
    try:
        if not _updates_enabled():
            raise BundleError("feature_disabled", "Partial bundle updates are not enabled")
        root = Path(params.directory).expanduser().resolve()
        _authorize(root)
        base = read_snapshot(root)
        if base is None or base.revision != params.base_revision:
            raise BundleError("revision_conflict", "Inspect the current bundle before updating",
                              base.revision if base else None)
        spec = DocumentBundleParams.model_validate(base.manifest["spec"]).model_dump()
        changes = params.changes.model_dump(exclude_unset=True)
        if not changes:
            raise BundleError("empty_changes", "Specify at least one change")
        patches = changes.pop("documents", [])
        if patches is None:
            raise BundleError("invalid_changes", "Document patches must be a list")
        by_name = {d["filename"]: d for d in spec["documents"]}
        seen: set[str] = set()
        for patch in patches:
            name = patch.pop("filename")
            if name not in by_name or name in seen:
                raise BundleError("invalid_changes", f"Unknown or duplicate document: {name}")
            seen.add(name)
            by_name[name].update(patch)
        spec.update(changes)
        spec["directory"] = str(root)
        updated = DocumentBundleParams.model_validate(spec)
        return await _create_version(updated, base, params.expected_source_sha256)
    except Exception as exc:
        return _failure(exc)


def register_document_bundle_capability(registry: CapabilityRegistry) -> None:
    registry.register(CapabilityDefinition(
        name="document_bundle",
        description=("Create/update consistent XLSX, DOCX, PPTX and PDF files from one CSV/XLSX source. "
                     "Specify row filters, aggregates and {{metric_id}} references in document content. "
                     "Reopens every file and publishes a complete version with source/output hashes. "
                     "Inspect the current version before changing an existing bundle."),
        domain=Domain.FILE, risk_level=RiskLevel.MEDIUM, group="write",
        parameters_model=DocumentBundleParams, execute=document_bundle,
    ))

    registry.register(CapabilityDefinition(
        name="document_bundle_inspect",
        description="Read a published bundle specification, source hashes, metrics and artifact integrity before updating it.",
        domain=Domain.FILE, risk_level=RiskLevel.LOW, group="read",
        parameters_model=BundleInspectParams, execute=document_bundle_inspect,
    ))
    if _updates_enabled():
        registry.register(CapabilityDefinition(
            name="document_bundle_update",
            description="Update selected fields of an inspected office bundle. Supply its base revision and current source SHA-256. Publishes all related files together, or reports a conflict without overwriting newer work.",
            domain=Domain.FILE, risk_level=RiskLevel.MEDIUM, group="write",
            parameters_model=BundleUpdateParams, execute=document_bundle_update,
        ))
