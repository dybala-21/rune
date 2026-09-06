"""Read tabular data, calculate aggregates, and resolve metric references."""

from __future__ import annotations

import csv
import io
import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field


class RowFilter(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str
    operator: Literal["eq", "ne", "in", "not_in"] = "eq"
    values: list[str | int | float] = Field(min_length=1)


class BundleMetric(BaseModel):
    model_config = ConfigDict(extra="forbid")
    id: str = Field(pattern=r"^[a-zA-Z][a-zA-Z0-9_]*$")
    operation: Literal["sum", "count", "mean", "min", "max"]
    column: str | None = None
    decimals: int = Field(default=0, ge=0, le=6)


def read_table(data: bytes, suffix: str, sheet: str | None) -> list[dict[str, Any]]:
    """Parse CSV or values-only XLSX bytes; reject formula cells."""
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            if sheet is None and len(wb.sheetnames) != 1:
                raise ValueError("Specify sheet for a workbook with multiple worksheets")
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = []
            for cells in ws.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Source contains formulas; supply a recalculated values-only snapshot")
                rows.append([cell.value for cell in cells])
                if len(rows) > 100_001:
                    raise ValueError("Source exceeds 100000 data rows")
        finally:
            wb.close()
    else:
        raise ValueError("Bundle source must be .csv or .xlsx")
    if not rows or len(rows) > 100_001:
        raise ValueError("Source must have a header and at most 100000 data rows")
    headers = rows[0]
    if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ValueError("Every source column needs a nonempty text header")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate source column headers")
    records = []
    for row in rows[1:]:
        if not any(v is not None and v != "" for v in row):
            continue
        if len(row) != len(headers):
            raise ValueError("Source row length does not match header")
        records.append(dict(zip(headers, row, strict=True)))
    if not records:
        raise ValueError("Source has no data rows")
    return records


def calculate(
    rows: list[dict[str, Any]], filters: list[RowFilter], metrics: list[BundleMetric],
) -> tuple[dict[str, int | float], dict[str, str], int]:
    columns = rows[0].keys()
    for f in filters:
        if f.column not in columns:
            raise ValueError(f"Unknown filter column: {f.column}")
        if f.operator in ("eq", "ne") and len(f.values) != 1:
            raise ValueError(f"{f.operator} requires exactly one value")
    for m in metrics:
        if m.operation != "count" and m.column not in columns:
            raise ValueError(f"Unknown metric column: {m.column}")
    selected = [r for r in rows if all(
        (r[f.column] in f.values) == (f.operator in ("eq", "in")) for f in filters
    )]
    values: dict[str, int | float] = {}
    display: dict[str, str] = {}
    for m in metrics:
        if m.id in values:
            raise ValueError(f"Duplicate metric ID: {m.id}")
        if m.operation == "count":
            value = Decimal(len(selected))
        else:
            if m.column is None:
                raise ValueError(f"Metric {m.id} requires a column")
            numbers = []
            for row in selected:
                raw = row[m.column]
                try:
                    number = Decimal(str(raw))
                except (InvalidOperation, ValueError) as exc:
                    raise ValueError(f"Non-numeric value in {m.column}: {raw!r}") from exc
                if not number.is_finite():
                    raise ValueError(f"Non-finite value in {m.column}")
                numbers.append(number)
            if m.operation != "sum" and not numbers:
                raise ValueError(f"{m.operation} is undefined for empty selection")
            if m.operation == "sum":
                value = sum(numbers, Decimal(0))
            elif m.operation == "mean":
                value = sum(numbers, Decimal(0)) / len(numbers)
            elif m.operation == "min":
                value = min(numbers)
            else:
                value = max(numbers)
        rounded = value.quantize(Decimal(10) ** -m.decimals, rounding=ROUND_HALF_UP)
        if len(rounded.as_tuple().digits) > 15:
            raise ValueError(f"Metric {m.id} exceeds Excel's 15-digit precision")
        values[m.id] = int(rounded) if rounded == rounded.to_integral() else float(rounded)
        display[m.id] = f"{rounded:,.{m.decimals}f}"
    return values, display, len(selected)


_REFERENCE = re.compile(r"\{\{([a-zA-Z][a-zA-Z0-9_]*)\}\}")


def resolve_references(
    spec: Any, values: dict[str, int | float], display: dict[str, str],
) -> Any:
    """A whole-cell reference stays numeric; references in prose use formatting."""
    if isinstance(spec, dict):
        return {k: resolve_references(v, values, display) for k, v in spec.items()}
    if isinstance(spec, list):
        return [resolve_references(v, values, display) for v in spec]
    if not isinstance(spec, str):
        return spec
    if match := _REFERENCE.fullmatch(spec):
        if match[1] not in values:
            raise ValueError(f"Unknown metric reference: {match[1]}")
        return values[match[1]]

    def substitute(match: re.Match[str]) -> str:
        if match[1] not in display:
            raise ValueError(f"Unknown metric reference: {match[1]}")
        return display[match[1]]

    result = _REFERENCE.sub(substitute, spec)
    if "{{" in result or "}}" in result:
        raise ValueError(f"Invalid metric reference: {spec}")
    return result
