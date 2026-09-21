"""Compare a saved table with aggregates computed from a fixed source snapshot."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation, localcontext
from itertools import pairwise
from typing import Annotated, Any, Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

from rune.capabilities.bundle_data import RowFilter


class Aggregate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(min_length=1, max_length=128)
    operation: Literal["sum", "count", "mean", "min", "max"]
    column: str | None = None
    decimals: int = Field(default=0, ge=0, le=6)


class TableSort(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str
    direction: Literal["asc", "desc"] = "asc"
    numeric: bool = False


class TablePlan(BaseModel):
    model_config = ConfigDict(extra="forbid")
    applicable: bool
    requirements: list[str] = Field(default_factory=list, max_length=24)
    filters: list[RowFilter] = Field(default_factory=list, max_length=20)
    deduplicate_by: list[str] = Field(default_factory=list, max_length=32)
    group_by: list[str] = Field(default_factory=list, max_length=4)
    grand_total: list[str] = Field(default_factory=list, max_length=4)
    exact_headers: bool = True
    exact_total_label: bool = True
    csv_bom: bool | None = None
    aggregates: list[Aggregate] = Field(default_factory=list, max_length=32)
    order_by: list[TableSort] = Field(default_factory=list, max_length=8)
    output_names: list[Annotated[str, Field(pattern=r"^[^/\\]+\.(?:[cC][sS][vV]|[xX][lL][sS][xX])$")]] = Field(
        default_factory=list, max_length=12, description="Output CSV/XLSX basenames only; no directories or surrounding prose.")
    output_sheet: str | None = Field(default=None, min_length=1, max_length=31)
    unverified: list[str] = Field(default_factory=list, max_length=24)
    out_of_scope: list[str] = Field(default_factory=list, max_length=24)

    @model_validator(mode="after")
    def valid_plan(self) -> TablePlan:
        if self.applicable and (not self.requirements or not self.aggregates):
            raise ValueError("An applicable plan needs quoted requirements and aggregates")
        if not self.applicable and (self.filters or self.deduplicate_by or self.group_by or self.aggregates or self.grand_total or self.order_by):
            raise ValueError("An inapplicable plan cannot contain executable conditions")
        if not self.applicable and not self.unverified:
            raise ValueError("Explain why the request cannot be checked")
        names = self.group_by + [a.name for a in self.aggregates]
        if self.grand_total and (len(self.grand_total) != len(self.group_by)
                                 or not any(label.strip() for label in self.grand_total)
                                 or any(len(label) > 128 for label in self.grand_total)):
            raise ValueError("Grand-total labels must identify each grouping column")
        if len(set(names)) != len(names) or len(set(self.deduplicate_by)) != len(self.deduplicate_by):
            raise ValueError("Duplicate output columns or duplicate keys")
        ordering = [rule.column for rule in self.order_by]
        if len(set(ordering)) != len(ordering) or set(ordering) - set(names):
            raise ValueError("Sort columns must uniquely identify output columns")
        for name in self.output_names:
            if "/" in name or "\\" in name or not name.lower().endswith((".csv", ".xlsx")):
                raise ValueError("Expected output names must be CSV/XLSX basenames")
        return self


def read_tabular(data: bytes, suffix: str, sheet: str | None = None,
                 header_row: int = 1) -> tuple[list[str], list[dict[str, Any]]]:
    if len(data) > 10_000_000:
        raise ValueError("Table exceeds 10 MB")
    if suffix == ".csv":
        if sheet is not None:
            raise ValueError("CSV does not have worksheets")
        stream = io.StringIO(data.decode("utf-8-sig"))
        iterator, close = csv.reader(stream), stream.close
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 64_000_000:
                raise ValueError("Expanded workbook exceeds 64 MB")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            if sheet is None and len(wb.sheetnames) != 1:
                raise ValueError("Specify the worksheet when a workbook has multiple sheets")
            ws = wb[sheet] if sheet else wb.worksheets[0]
            if ws.max_column and ws.max_column > 256:
                raise ValueError("Table exceeds 256 columns")

            def values():
                for cells in ws.iter_rows():
                    if any(c.data_type in {"f", "e"} for c in cells):
                        raise ValueError("Formula/error cells require a recalculated values-only snapshot")
                    yield [c.value for c in cells]

            iterator, close = values(), wb.close
        except Exception:
            wb.close()
            raise
    else:
        raise ValueError("Only CSV and values-only XLSX are supported")
    try:
        headers: list[str] = []
        records = []
        cells_seen = 0
        for number, row in enumerate(iterator, 1):
            cells_seen += len(row)
            if number > 100_100 or cells_seen > 2_000_000 or len(row) > 256:
                raise ValueError("Table exceeds the row or cell limit")
            if number < header_row:
                continue
            if number == header_row:
                if not row or any(not isinstance(v, str) or not v.strip() for v in row):
                    raise ValueError("Each column needs a nonempty text header")
                if len(set(row)) != len(row):
                    raise ValueError("Duplicate column headers")
                headers = list(row)
                continue
            if not any(v is not None and v != "" for v in row):
                continue
            if len(row) != len(headers):
                raise ValueError("Row width differs from the header")
            records.append(dict(zip(headers, row, strict=True)))
            if len(records) > 100_000:
                raise ValueError("Table exceeds 100000 data rows")
        if not headers:
            raise ValueError("Header row was not found")
        return headers, records
    finally:
        close()


def number(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Expected a number, got {value!r}")
    text = str(value).strip()
    if "," in text:
        # Validate grouping before removing separators; 12,34 must not become 1234.
        if not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?:[eE][+-]?\d+)?", text):
            raise ValueError(f"Invalid numeric grouping: {value!r}; expected comma groups of three and a decimal point")
        text = text.replace(",", "")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Expected a number, got {value!r}") from exc
    if not result.is_finite() or abs(result.adjusted()) > 100 or len(result.as_tuple().digits) > 100:
        raise ValueError("Numeric value is non-finite or exceeds supported precision")
    return result


def expected_table(plan: TablePlan, headers: list[str], rows: list[dict[str, Any]]
                   ) -> tuple[list[tuple[str, ...]], dict[str, int]]:
    required = set(plan.group_by + plan.deduplicate_by)
    required.update(f.column for f in plan.filters)
    required.update(a.column for a in plan.aggregates if a.operation != "count")
    if required - set(headers):
        raise ValueError(f"Unknown source columns: {required - set(headers)}")
    for f in plan.filters:
        if f.operator in {"eq", "ne"} and len(f.values) != 1:
            raise ValueError(f"{f.operator} needs exactly one value")
    selected = [row for row in rows if all(
        (row[f.column] in f.values) == (f.operator in {"eq", "in"}) for f in plan.filters
    )]
    filtered = len(rows) - len(selected)
    duplicates = 0
    if plan.deduplicate_by:
        unique = {}
        for row in selected:
            key = tuple(row[c] for c in plan.deduplicate_by)
            if any(v is None or v == "" for v in key):
                raise ValueError("Duplicate keys contain empty values; clarify the missing-key policy")
            if key in unique:
                if row != unique[key]:
                    raise ValueError(f"Conflicting rows for duplicate key {key!r}; clarify which record to keep")
                duplicates += 1
            else:
                unique[key] = row
        selected = list(unique.values())
    groups: dict[tuple[str, ...], list[dict[str, Any]]] = {} if plan.group_by else {(): []}
    for row in selected:
        key = tuple("" if row[c] is None else str(row[c]) for c in plan.group_by)
        groups.setdefault(key, []).append(row)
        if len(groups) > 10_000:
            raise ValueError("Result exceeds 10000 groups")
    if plan.grand_total:
        key = tuple(plan.grand_total)
        if key in groups:
            raise ValueError("Grand-total label collides with a source group; choose an unambiguous label")
        # Recompute from source rows, before group rounding or averaging.
        groups[key] = selected
    expected = []
    with localcontext() as context:
        context.prec = 220
        for key, members in groups.items():
            values = list(key)
            for aggregate in plan.aggregates:
                numbers = ([number(row[aggregate.column]) for row in members]
                           if aggregate.operation != "count" else [])
                if not numbers and aggregate.operation not in {"count", "sum"}:
                    raise ValueError(f"{aggregate.operation} is undefined for an empty selection")
                match aggregate.operation:
                    case "count":
                        value = Decimal(len(members))
                    case "sum":
                        value = sum(numbers, Decimal(0))
                    case "mean":
                        value = sum(numbers, Decimal(0)) / len(numbers)
                    case "min":
                        value = min(numbers)
                    case "max":
                        value = max(numbers)
                value = value.quantize(Decimal(10) ** -aggregate.decimals, rounding=ROUND_HALF_UP)
                values.append(str(value))
            expected.append(tuple(values))
    return expected, {"source_rows": len(rows), "filtered_rows": filtered,
                      "duplicates_removed": duplicates, "selected_rows": len(selected),
                      "output_rows": len(expected)}


def compare_table(plan: TablePlan, expected: list[tuple[str, ...]],
                  headers: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    columns = plan.group_by + [a.name for a in plan.aggregates]
    if len(headers) != len(columns) or plan.exact_headers and headers != columns:
        return {"status": "fail", "issues": [{"check": "columns", "expected": columns, "actual": headers}]}

    def canonical(row: tuple[str, ...]) -> tuple[str | Decimal, ...]:
        return tuple(row[:len(plan.group_by)]) + tuple(number(v) for v in row[len(plan.group_by):])

    actual = [tuple("" if row[c] is None else str(row[c]) for c in headers) for row in rows]
    if plan.grand_total and not plan.exact_total_label:
        width = len(plan.group_by)
        group_keys = {row[:width] for row in expected[:-1]}
        # An extra row may represent the requested total, but its values and multiplicity must still match.
        actual = [tuple(plan.grand_total) + row[width:] if row[:width] not in group_keys else row for row in actual]
    wanted, found = Counter(map(canonical, expected)), Counter(map(canonical, actual))
    missing, extra = wanted - found, found - wanted
    issues = []
    for check, difference in (("missing_or_incorrect_rows", missing), ("unexpected_or_duplicate_rows", extra)):
        if difference:
            issues.append({"check": check, "count": sum(difference.values()),
                           "examples": [[str(v) for v in row] for row in list(difference)[:5]]})
    if plan.order_by:
        ordered = (row for row in map(canonical, actual)
                   if not plan.grand_total or row[:len(plan.group_by)] != tuple(plan.grand_total))
        rules = [(columns.index(rule.column), rule) for rule in plan.order_by]
        for left, right in pairwise(ordered):
            out_of_order = False
            for index, rule in rules:
                a, b = (number(left[index]), number(right[index])) if rule.numeric else (left[index], right[index])
                if a == b:
                    continue
                out_of_order = a > b if rule.direction == "asc" else a < b
                if out_of_order:
                    issues.append({"check": "row_order", "column": rule.column, "direction": rule.direction})
                break
            if out_of_order:
                break
    return {"status": "fail" if issues else "pass", "issues": issues,
            "checks": ["columns", "aggregates", "group_coverage", "row_multiplicity", *(["row_order"] if plan.order_by else [])]}
